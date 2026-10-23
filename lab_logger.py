#!/usr/bin/env python3
"""
lab_logger.py - Automatic data logger for bench instruments.

Supported instruments
  * OHAUS MB90 moisture analyzer      (serial over USB or RS232)
  * OHAUS pH meter a-AB33 / a-AB41    (serial over USB or RS232)
  * Linshang LS177C coffee colorimeter (watches a folder for Excel/CSV exports
                                        made with Linshang's PC software)
  * AQUALAB 3 water activity meter     (pulls readings straight from SKALA Control
                                        when online, type "skala_api"; or watches a
                                        folder of downloaded readings, "export_folder")

Everything is written as CSV files into OUTPUT_FOLDER. Point OUTPUT_FOLDER at a
OneDrive / Google Drive / Dropbox folder and the data is backed up to the cloud
automatically.

Usage
  python lab_logger.py                  start logging (creates config on first run)
  python lab_logger.py --list-ports     show the serial ports this computer can see
  python lab_logger.py --monitor COM5   show raw text arriving on one port (testing)
  python lab_logger.py --selftest       run the parsers on sample printouts
  python lab_logger.py --demo           open the dashboard with simulated instruments

While running, the dashboard is at http://localhost:8080 (this laptop only).

While running, type these commands in the window and press Enter:
  lot <number>     set the lot/batch number added to every new record
  op <name>        set the operator name added to every new record
  note <text>      set a free-text note added to every new record
  sync             check SKALA Control for new AQUALAB 3 readings now
  status           show connection status of each instrument
  quit             stop the logger

Requires:  pip install pyserial openpyxl
"""

import argparse
import csv
import datetime as dt
import json
import os
import re
import shutil
import sys
import threading
import time
from pathlib import Path

try:
    import serial
    import serial.tools.list_ports
except ImportError:  # allow --selftest without pyserial
    serial = None

try:
    import openpyxl
except ImportError:
    openpyxl = None


SCRIPT_DIR = Path(__file__).resolve().parent
CONFIG_PATH = SCRIPT_DIR / "lab_logger_config.json"

DEFAULT_CONFIG = {
    "_help": "Edit this file, save it, then restart lab_logger.py. "
             "Use forward slashes or double backslashes in Windows paths.",
    "station_name": "Lab bench",
    "output_folder": (SCRIPT_DIR / "data").as_posix(),
    "shared_file": {"enabled": True, "aqualab_lot_from": "batch_name"},
    "dashboard": {"enabled": True, "port": 8080, "open_browser": True,
                  "support_contact": "your supervisor"},
    "upload": {"folder": ""},
    "history": {"enabled": True, "keep_days": 0},
    "lot_rules": {"expire_minutes": 120, "clear_at_midnight": True, "lot_pattern": "", "lot_example": "",
                  "operators": [], "block_upload_without_lot": True},
    "instruments": [
        {
            "name": "MB90",
            "type": "mb90",
            "enabled": True,
            "port": "COM6",
            "baudrate": 9600,
            "bytesize": 8,
            "parity": "N",
            "stopbits": 1,
            "xonxoff": False,
        },
        {
            "name": "pH_meter",
            "type": "ph",
            "enabled": True,
            "port": "COM3",
            "baudrate": 9600,
            "bytesize": 8,
            "parity": "N",
            "stopbits": 1,
            "xonxoff": False,
        },
        {
            "name": "LS177C",
            "type": "colorimeter_folder",
            "enabled": True,
            "watch_folder": (SCRIPT_DIR / "data" / "Colorimeter Exports").as_posix(),
            "poll_seconds": 10,
        },
        {
            "name": "AQUALAB3",
            "type": "skala_api",
            "enabled": True,
            "token_file": (SCRIPT_DIR / "skala_token.txt").as_posix(),
            "sbs_serial": "SKM16F64",
            "poll_minutes": 5,
            "lookback_days": 3,
        },
    ],
}

# Shared, operator-set context added to every record.
CONTEXT = {"lot": "", "operator": "", "note": "", "set_at": ""}
CONTEXT_LOCK = threading.Lock()
# Rules that stop an old lot or name being reused by accident (set in the config, "lot_rules")
LOT_RULES = {"expire_minutes": 120, "clear_at_midnight": True, "lot_pattern": "", "lot_example": "",
             "operators": [], "block_upload_without_lot": True}
STATUS = {}
STOP = threading.Event()

# Live state for the dashboard (localhost web page)
from collections import deque
DEVICES = {}          # name -> {"type","label","ok","detail","last_data","last_result"}
LIVE = {}             # name -> {"lines": deque, "test": {...}}
RECENT = deque(maxlen=300)   # newest shared results, for the dashboard
DEV_LOCK = threading.Lock()
DEVICE_LABELS = {"mb90": "MB90 moisture", "ph": "pH meter",
                 "colorimeter_folder": "LS177C color", "export_folder": "AQUALAB 3",
                 "skala_api": "AQUALAB 3"}


def set_status(name, text, ok=None, **extra):
    """Update the console status text and the dashboard's device record."""
    STATUS[name] = text
    with DEV_LOCK:
        d = DEVICES.setdefault(name, {"name": name, "ok": False, "detail": "", "last_data": "",
                                      "last_result": None, "type": "", "label": name})
        d["detail"] = text
        if ok is not None:
            d["ok"] = bool(ok)
        d.update(extra)


def live_line(name, line):
    with DEV_LOCK:
        lv = LIVE.setdefault(name, {"lines": deque(maxlen=200), "test": None})
        lv["lines"].append([dt.datetime.now().strftime("%H:%M:%S"), line])
        if name in DEVICES:
            DEVICES[name]["last_data"] = now_iso()


# --------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------
def now_iso():
    return dt.datetime.now().isoformat(timespec="seconds")


LOG_FILE = SCRIPT_DIR / "lab_logger.log"


def write_log_file(msg):
    try:
        with open(LOG_FILE, "a", encoding="utf-8") as f:
            f.write(f"{dt.datetime.now():%Y-%m-%d %H:%M:%S} {msg}\n")
    except Exception:
        pass


def log(msg):
    line = f"[{dt.datetime.now():%H:%M:%S}] {msg}"
    if sys.stdout is not None:
        try:
            print(line, flush=True)
            return
        except Exception:
            pass
    try:  # started with no window (pythonw / desktop launcher): write to a log file
        with open(LOG_FILE, "a", encoding="utf-8") as f:
            f.write(f"{dt.date.today()} {line}\n")
    except Exception:
        pass


def expand(path_str):
    return Path(os.path.expandvars(os.path.expanduser(path_str)))


def context_snapshot():
    with CONTEXT_LOCK:
        return dict(CONTEXT)


class CsvWriter:
    """Appends rows to CSV files. If a file is locked (e.g. open in Excel),
    rows are held in memory and retried, so nothing is lost."""

    def __init__(self):
        self.lock = threading.Lock()
        self.pending = []  # list of (path, rows)

    def append(self, path, rows):
        if not rows:
            return
        with self.lock:
            self.pending.append((Path(path), rows))
            self._flush_locked()

    def flush(self):
        with self.lock:
            self._flush_locked()

    def _flush_locked(self):
        still_pending = []
        for path, rows in self.pending:
            try:
                self._write(path, rows)
            except PermissionError:
                still_pending.append((path, rows))
        if still_pending and len(still_pending) != getattr(self, "_last_warn", 0):
            log(f"WARNING: {still_pending[0][0].name} is locked (open in Excel?). "
                f"Close it; {sum(len(r) for _, r in still_pending)} row(s) waiting.")
        self._last_warn = len(still_pending)
        self.pending = still_pending

    @staticmethod
    def _write(path, rows):
        path.parent.mkdir(parents=True, exist_ok=True)
        new_fields = []
        for r in rows:
            for k in r:
                if k not in new_fields:
                    new_fields.append(k)

        if path.exists() and path.stat().st_size > 0:
            with open(path, newline="", encoding="utf-8-sig") as f:
                reader = csv.reader(f)
                header = next(reader, [])
            missing = [k for k in new_fields if k not in header]
            if missing:
                # Header grew: rewrite the file with the combined header.
                with open(path, newline="", encoding="utf-8-sig") as f:
                    old_rows = list(csv.DictReader(f))
                header = header + missing
                tmp = path.with_suffix(".tmp")
                with open(tmp, "w", newline="", encoding="utf-8-sig") as f:
                    w = csv.DictWriter(f, fieldnames=header)
                    w.writeheader()
                    w.writerows(old_rows)
                    w.writerows(rows)
                os.replace(tmp, path)
                return
            with open(path, "a", newline="", encoding="utf-8-sig") as f:
                w = csv.DictWriter(f, fieldnames=header, extrasaction="ignore")
                w.writerows(rows)
        else:
            with open(path, "w", newline="", encoding="utf-8-sig") as f:
                w = csv.DictWriter(f, fieldnames=new_fields)
                w.writeheader()
                w.writerows(rows)


WRITER = CsvWriter()


def monthly_csv(out_dir, name):
    return out_dir / name / f"{name}_{dt.date.today():%Y-%m}.csv"


def write_raw(out_dir, name, text):
    raw_dir = out_dir / name / "raw"
    raw_dir.mkdir(parents=True, exist_ok=True)
    with open(raw_dir / f"{name}_raw_{dt.date.today():%Y-%m-%d}.txt", "a",
              encoding="utf-8") as f:
        f.write(f"===== received {now_iso()} =====\n{text}\n\n")


def first_float(pattern, text, flags=re.I):
    m = re.search(pattern, text, flags)
    return m.group(1) if m else ""


# --------------------------------------------------------------------------
# Parsers
# --------------------------------------------------------------------------
MB90_LABELS = [
    ("serial_number", r"SNR\s*\(Drying Unit\)\s*(.+)"),
    ("software_version", r"SW\s*\(Drying Unit\)\s*(.+)"),
    ("method_name", r"Method Name\s+(.+)"),
    ("drying_program", r"Drying Prog\w*\s+(.+)"),
    ("drying_temp_C", r"Drying Temp\w*\s+(-?[\d.]+)"),
    ("switch_off", r"Switch Off\s+(.+)"),
    ("start_weight_g", r"Start Weight\s+(-?[\d.]+)"),
    ("total_time", r"Total Time\.?\s+([\d:]+)"),
    ("end_result", r"End Result\s*\.?\s+(-?[\d.]+)"),
    ("end_result_unit", r"End Result\s*\.?\s+-?[\d.]+\s*(%MC|%DC|%RG|g)"),
    ("nominal_weight_g", r"Nominal Weight\s+(-?[\d.]+)"),
    ("actual_weight_g", r"Actual Weight\s+(-?[\d.]+)"),
    ("difference_g", r"Difference\s+(-?[\d.]+)"),
    ("cell_temperature_C", r"Cell Temperature\s+(-?[\d.]+)"),
    ("temp1_target_C", r"Temp1 target\s+(-?[\d.]+)"),
    ("temp1_actual_C", r"Temp1 actual\s+(-?[\d.]+)"),
    ("temp2_target_C", r"Temp2 target\s+(-?[\d.]+)"),
    ("temp2_actual_C", r"Temp2 actual\s+(-?[\d.]+)"),
    ("adjustment", r"Adjustment\s+(Done|Failed|\w+)"),
    ("stat_sample_number", r"Sample Number\s+(\d+)"),
    ("stat_mean", r"Mean Value\s+(-?[\d.]+)"),
    ("stat_std_dev", r"Standard Deviation\s+(-?[\d.]+)"),
    ("stat_min", r"Minimum Value\s+(-?[\d.]+)"),
    ("stat_max", r"Maximum Value\s+(-?[\d.]+)"),
]
MB90_CURVE = re.compile(r"^\s*(\d{1,3}:\d{2})\s*min\s+(-?[\d.:]+)\s*(%MC|%DC|%RG|g)", re.I)


def parse_mb90(text):
    upper = text.upper()
    if "WEIGHT ADJUST" in upper:
        rtype = "weight_adjustment"
    elif "TEMPERATURE ADJUST" in upper:
        rtype = "temperature_adjustment"
    elif "STATISTICS" in upper:
        rtype = "statistics"
    elif "MOISTURE DETERMINATION" in upper:
        rtype = "moisture_test"
    else:
        rtype = "other"

    rec = {"record_type": rtype}
    for key, pat in MB90_LABELS:
        rec[key] = ""
        for line in text.splitlines():
            m = re.search(pat, line, re.I)
            if m:
                rec[key] = m.group(1).strip()
                break

    curve = []
    for line in text.splitlines():
        m = MB90_CURVE.match(line)
        if m:
            # "29:36" style typos in printouts -> treat ':' as '.'
            curve.append((m.group(1), m.group(2).replace(":", "."), m.group(3)))
    rec["curve_points"] = len(curve)
    if rtype == "moisture_test" and not rec["end_result"] and curve:
        rec["end_result"], rec["end_result_unit"] = curve[-1][1], curve[-1][2]
        rec["record_type"] = "moisture_test_incomplete"
    return rec, curve


def parse_ph(text):
    flat = " ".join(text.split())
    rec = {
        "meter_model": first_float(r"\b(A?-?AB\d{2}[A-Z0-9]+)\b", flat),
        "meter_datetime": first_float(r"(\d{2,4}/\d{1,2}/\d{1,4}\s+\d{1,2}:\d{2}(?::\d{2})?)", flat),
        "user_id": first_float(r"User\s*(?:ID)?\s*[:#]?\s*(\d+)", flat),
        "sample_id": first_float(r"Sample\s*(?:ID)?\s*[:#]?\s*(\w+)", flat),
        "mode": first_float(r"\b(ORP|mV|pH|Cond|TDS|SALT|RES)\b(?=\s+-?[\d.]+)", flat, 0),
        "pH": first_float(r"(-?\d+\.\d+)\s*pH\b", flat),
        "mV": first_float(r"(-?\d+\.\d+)\s*m[Vv]\b", flat),
        "ORP_RmV": first_float(r"(-?\d+\.\d+)\s*RmV\b", flat),
        "temperature": first_float(r"(-?\d+\.\d+)\s*°?\s*[℃℉CF]°?(?![A-Za-z/])", flat, 0),
        "temp_unit": first_float(r"-?\d+\.\d+\s*°?\s*([CF℃℉])°?(?![A-Za-z/])", flat, 0)
                     .replace("℃", "C").replace("℉", "F"),
        "temp_type": first_float(r"\b(ATC|MTC)\b", flat),
        "slope_pct": first_float(r"(-?\d+\.\d+)\s*%(?!/)", flat),
        "offset_mV": "",
        "conductivity": first_float(r"(-?\d+\.?\d*)\s*[uµm]S/cm", flat),
        "tds_mg_L": first_float(r"(-?\d+\.?\d*)\s*mg/L", flat),
        "calibration_record": "yes" if re.search(r"Cal\s*Point|Calibration|Offset\s*mV", flat, re.I) else "",
    }
    mvs = re.findall(r"(-?\d+\.\d+)\s*m[Vv]\b", flat)
    if len(mvs) >= 2:        # measurement printout: first mV is reading, last is offset
        rec["offset_mV"] = mvs[-1]
    sn = re.search(r"\b(\d{8,12})\b", flat)
    rec["serial_number"] = sn.group(1) if sn else ""
    return rec


# --------------------------------------------------------------------------
# Serial instrument worker
# --------------------------------------------------------------------------
class SerialInstrument(threading.Thread):
    START_MARKERS = {"ph": re.compile(r"\bA?-?AB(33|41)", re.I)}
    END_MARKERS = {"mb90": re.compile(r"-{3,}\s*END\s*-{3,}", re.I)}
    IDLE_FLUSH = {"mb90": 900.0, "ph": 2.5}   # seconds of silence before a record is closed

    def __init__(self, cfg, out_dir):
        super().__init__(daemon=True)
        self.cfg = cfg
        self.name_ = cfg["name"]
        self.kind = cfg["type"]
        self.out_dir = out_dir
        self.buffer = []
        self.last_rx = 0.0
        set_status(self.name_, "starting", False, type=self.kind,
                   label=cfg.get("label") or DEVICE_LABELS.get(self.kind, self.name_))

    def open_port(self):
        parity = {"N": serial.PARITY_NONE, "E": serial.PARITY_EVEN,
                  "O": serial.PARITY_ODD}[self.cfg.get("parity", "N").upper()[0]]
        return serial.Serial(
            port=self.cfg["port"],
            baudrate=int(self.cfg.get("baudrate", 9600)),
            bytesize=int(self.cfg.get("bytesize", 8)),
            parity=parity,
            stopbits=float(self.cfg.get("stopbits", 1)),
            xonxoff=bool(self.cfg.get("xonxoff", False)),
            timeout=0.5,
        )

    def run(self):
        while not STOP.is_set():
            try:
                with self.open_port() as port:
                    set_status(self.name_, f"USB · {self.cfg['port']}", True)
                    log(f"{self.name_}: connected on {self.cfg['port']}")
                    partial = b""
                    while not STOP.is_set():
                        chunk = port.read(512)
                        if chunk:
                            self.last_rx = time.time()
                            partial += chunk
                            while b"\n" in partial or b"\r" in partial:
                                line, partial = re.split(rb"\r\n|\n|\r", partial, maxsplit=1)
                                self.handle_line(self.decode(line))
                        elif partial and time.time() - self.last_rx > 1.0:
                            self.handle_line(self.decode(partial))
                            partial = b""
                        self.check_idle()
            except Exception as e:  # port missing, unplugged, busy...
                if STATUS.get(self.name_, "").startswith("USB"):
                    log(f"{self.name_}: connection lost ({e}). Retrying...")
                elif not STATUS.get(self.name_, "").startswith("Not connected"):
                    log(f"{self.name_}: cannot open {self.cfg['port']} ({e}). "
                        f"Will keep retrying every 5 s.")
                set_status(self.name_, f"Not connected ({self.cfg['port']})", False)
                self.flush()
                STOP.wait(5)

    @staticmethod
    def decode(b):
        try:
            return b.decode("utf-8")
        except UnicodeDecodeError:
            return b.decode("latin-1")

    def handle_line(self, line):
        if line.strip():
            live_line(self.name_, line.rstrip())
            if self.kind == "mb90":
                self.track_mb90(line)
        if not line.strip():
            if self.buffer:
                self.buffer.append("")
            return
        start = self.START_MARKERS.get(self.kind)
        if start and start.search(line) and any(l.strip() for l in self.buffer):
            self.flush()
        self.buffer.append(line.rstrip())
        end = self.END_MARKERS.get(self.kind)
        if end and end.search(line):
            self.flush()

    def track_mb90(self, line):
        """Keep a live view of the MB90 test that is printing right now."""
        with DEV_LOCK:
            lv = LIVE.setdefault(self.name_, {"lines": deque(maxlen=200), "test": None})
            t = lv["test"]
            if re.search(r"MOISTURE DETERMINATION", line, re.I):
                lv["test"] = {"active": True, "started": now_iso(), "points": [], "method": "",
                              "temp": "", "start_weight": "", "value": "", "unit": "", "elapsed": ""}
                return
            if not t:
                return
            for key, pat in (("method", r"Method Name\s+(.+)"), ("temp", r"Drying Temp\w*\s+(-?[\d.]+)"),
                             ("start_weight", r"Start Weight\s+(-?[\d.]+)")):
                m = re.search(pat, line, re.I)
                if m:
                    t[key] = m.group(1).strip()
            m = MB90_CURVE.match(line)
            if m:
                val = m.group(2).replace(":", ".")
                t["points"].append([m.group(1), val])
                t.update(elapsed=m.group(1), value=val, unit=m.group(3))
            m = re.search(r"End Result\s*\.?\s+(-?[\d.]+)\s*(%MC|%DC|%RG|g)?", line, re.I)
            if m:
                t.update(value=m.group(1), unit=m.group(2) or t.get("unit", ""))
            if self.END_MARKERS["mb90"].search(line):
                t["active"] = False
                t["finished"] = now_iso()

    def check_idle(self):
        if self.buffer and time.time() - self.last_rx > self.IDLE_FLUSH.get(self.kind, 3.0):
            self.flush()

    def flush(self):
        text = "\n".join(self.buffer).strip()
        self.buffer = []
        if text:
            save_serial_record(self.name_, self.kind, text, self.out_dir)


# --------------------------------------------------------------------------
# Shared results file: All Results (one row per result) + Lot Summary
# --------------------------------------------------------------------------
SHARED_COLUMNS = [
    # all instruments
    "logged_at", "record_id", "instrument", "lot", "operator", "note",
    "measurement", "value", "unit", "raw_text",
    # MB90
    "record_type", "end_result", "end_result_unit", "start_weight_g", "total_time",
    "method_name", "drying_program", "drying_temp_C", "mb90_serial_number",
    "nominal_weight_g", "actual_weight_g", "difference_g", "cell_temperature_C", "adjustment",
    # pH meter
    "pH", "mV", "temperature", "temp_unit", "slope_pct", "offset_mV", "sample_id", "user_id",
    "meter_datetime", "meter_model", "ph_serial_number", "calibration_record",
    # LS177C
    "L*", "a*", "b*", "test_name", "colorimeter_datetime", "source_file",
    # AQUALAB 3
    "water_activity", "temperature_C", "duration_s", "reading_datetime", "aw_in_spec",
    "in_haccp_spec", "spec_min", "spec_max", "status", "product", "batch_name", "batch_date",
    "sample_fields", "dual_block_difference", "is_fast", "instrument_sn",
    "instrument_verified_at", "skala_user_id", "skala_reading_id",
]
# Lot Summary: (column title, measurement name in All Results, field to average).
# When a lot has several results for the same measurement, the summary shows the average.
SUMMARY_RESULTS = [
    ("Moisture (%MC)", "Moisture", "value"),
    ("pH", "pH", "value"),
    ("L*", "Color", "L*"),
    ("a*", "Color", "a*"),
    ("b*", "Color", "b*"),
    ("Water activity (aw)", "Water activity", "value"),
]
MB90_MEASURES = {"%MC": "Moisture", "%DC": "Solids", "%RG": "Regain", "g": "Final weight"}


def _num(v):
    try:
        return float(str(v).strip())
    except (TypeError, ValueError):
        return None


def _pick(row, *candidates):
    """Find a value in a colorimeter export row whose header matches one of the
    candidate names (case/space-insensitive). Used because the exact Linshang
    export headers aren't confirmed yet."""
    norm = {re.sub(r"[\s_]", "", str(k)).lower(): v for k, v in row.items()}
    for c in candidates:
        v = norm.get(re.sub(r"[\s_]", "", c).lower())
        if v not in (None, ""):
            return v
    return ""


class SharedFile:
    def __init__(self):
        self.folder = None
        self.columns = SHARED_COLUMNS
        self.aqualab_lot_from = "batch_name"
        self.enabled = False
        self.dirty = False
        self.lock = threading.Lock()

    def configure(self, cfg, out_dir):
        sc = cfg.get("shared_file", {})
        self.enabled = sc.get("enabled", True)
        self.folder = expand(sc["folder"]) if sc.get("folder") else out_dir / "Shared"
        self.columns = sc.get("columns") or SHARED_COLUMNS
        self.aqualab_lot_from = sc.get("aqualab_lot_from", "batch_name")
        if self.enabled:
            self.folder.mkdir(parents=True, exist_ok=True)
            log(f"Shared file: {self.folder}")

    def all_results_path(self):
        return self.folder / f"All_Results_{dt.date.today():%Y-%m}.csv"

    def add(self, rows):
        if not self.enabled or not rows:
            return
        out = [{c: r.get(c, "") for c in self.columns} for r in rows]
        WRITER.append(self.all_results_path(), out)
        self.dirty = True
        EXCEL.mark()
        for r in rows:
            remember_result(r)

    # ---- builders: map each instrument's record to shared columns ----
    @staticmethod
    def from_mb90(base, rec, raw):
        rtype = rec.get("record_type", "")
        r = {**base, **rec, "raw_text": raw, "mb90_serial_number": rec.get("serial_number", "")}
        if rtype.startswith("moisture_test"):
            unit = rec.get("end_result_unit") or "%MC"
            r.update(measurement=MB90_MEASURES.get(unit, "Moisture"), value=rec.get("end_result"), unit=unit)
        elif rtype == "weight_adjustment":
            r.update(measurement="MB90 weight calibration", value=rec.get("difference_g"), unit="g difference")
        elif rtype == "temperature_adjustment":
            r.update(measurement="MB90 temperature calibration", value=rec.get("adjustment"), unit="")
        elif rtype == "statistics":
            r.update(measurement="MB90 statistics", value=rec.get("stat_mean"), unit="mean")
        else:
            r.update(measurement="MB90 printout", value="", unit="")
        return r

    @staticmethod
    def from_ph(base, rec, raw):
        r = {**base, **rec, "raw_text": raw, "ph_serial_number": rec.get("serial_number", "")}
        if rec.get("calibration_record"):
            r.update(measurement="pH calibration", value=rec.get("slope_pct"), unit="% slope")
        else:
            r.update(measurement="pH", value=rec.get("pH"), unit="pH")
        return r

    @staticmethod
    def from_colorimeter(row):
        Lv = _pick(row, "Sample L*", "SampleL*", "L*", "L")
        av = _pick(row, "Sample a*", "Samplea*", "a*", "a")
        bv = _pick(row, "Sample b*", "Sampleb*", "b*", "b")
        meta = {"logged_at", "instrument", "lot", "operator", "note", "source_file", "sheet", "row"}
        raw = json.dumps({k: v for k, v in row.items() if k not in meta}, default=str)
        return {**row, "record_id": f"{row.get('instrument')}-{row.get('source_file')}-{row.get('row')}",
                "measurement": "Color", "value": Lv, "unit": "L*", "L*": Lv, "a*": av, "b*": bv,
                "std_L*": _pick(row, "Standard L*"), "std_a*": _pick(row, "Standard a*"),
                "std_b*": _pick(row, "Standard b*"),
                "test_name": _pick(row, "Test Name", "TestName", "Name"),
                "colorimeter_datetime": _pick(row, "DateTime", "Date Time", "Date", "Time"),
                "raw_text": raw}

    def from_aqualab(self, srow):
        r = dict(srow)
        r.update(record_id=f"{srow.get('instrument')}-{srow.get('reading_id')}",
                 measurement="Water activity", value=srow.get("water_activity"), unit="aw",
                 duration_s=srow.get("duration"), aw_in_spec=srow.get("in_spec"),
                 skala_reading_id=srow.get("reading_id"), raw_text=srow.get("raw_json", ""))
        if self.aqualab_lot_from == "batch_name" and srow.get("batch_name"):
            r["lot"] = srow.get("batch_name")
        return r

    def from_aqualab_export(self, row):
        aw = _pick(row, "Water Activity", "aw", "Aw", "Water Activity (aw)")
        r = {**row, "record_id": f"{row.get('instrument')}-{row.get('source_file')}-{row.get('row')}",
             "measurement": "Water activity", "value": aw, "unit": "aw", "water_activity": aw,
             "temperature_C": _pick(row, "Temperature", "Temp", "Temperature (C)"),
             "batch_name": _pick(row, "Batch", "Batch Name"), "product": _pick(row, "Product"),
             "reading_datetime": _pick(row, "Date", "DateTime", "Date Time")}
        if self.aqualab_lot_from == "batch_name" and r["batch_name"]:
            r["lot"] = r["batch_name"]
        return r

    # ---- Lot Summary ----
    def rebuild_summary(self):
        if not self.enabled or not self.dirty:
            return
        with self.lock:
            fields, out = summarize_rows(read_all_results())
            path = self.folder / "Lot_Summary.csv"
            tmp = path.with_suffix(".tmp")
            try:
                with open(tmp, "w", newline="", encoding="utf-8-sig") as fh:
                    w = csv.DictWriter(fh, fieldnames=fields, extrasaction="ignore")
                    w.writeheader()
                    w.writerows(out)
                os.replace(tmp, path)
                self.dirty = False
            except PermissionError:
                pass  # Lot_Summary.csv open in Excel; retried every few seconds


SHARED = SharedFile()


def remember_result(r):
    item = {"time": r.get("reading_datetime") or r.get("logged_at") or "",
            "logged_at": r.get("logged_at", ""), "lot": r.get("lot", ""), "record_id": r.get("record_id", ""),
            "instrument": r.get("instrument", ""), "measurement": r.get("measurement", ""),
            "value": r.get("value", ""), "unit": r.get("unit", ""), "operator": r.get("operator", "")}
    with DEV_LOCK:
        RECENT.appendleft(item)
        d = DEVICES.get(item["instrument"])
        if d is not None:
            d["last_result"] = item


def preload_recent():
    """Fill the dashboard's recent list from this month's shared file at start-up."""
    if not SHARED.enabled:
        return
    path = SHARED.all_results_path()
    if not path.exists():
        return
    try:
        with open(path, newline="", encoding="utf-8-sig") as fh:
            rows = list(csv.DictReader(fh))
    except Exception:
        return
    for r in rows[-RECENT.maxlen:]:
        remember_result(r)


def save_serial_record(name, kind, text, out_dir):
    ctx = context_snapshot()
    logged = now_iso()
    record_id = f"{name}-{dt.datetime.now():%Y%m%d-%H%M%S-%f}"[:-3]
    base = {"logged_at": logged, "record_id": record_id, "instrument": name,
            "lot": ctx["lot"], "operator": ctx["operator"], "note": ctx["note"]}
    write_raw(out_dir, name, text)

    if kind == "mb90":
        rec, curve = parse_mb90(text)
        row = {**base, **rec, "raw_text": text.replace("\n", " | ")}
        WRITER.append(monthly_csv(out_dir, name), [row])
        if curve:
            WRITER.append(out_dir / name / f"{name}_curves_{dt.date.today():%Y-%m}.csv",
                          [{"record_id": record_id, "logged_at": logged, "lot": ctx["lot"],
                            "elapsed": t, "value": v, "unit": u} for t, v, u in curve])
        SHARED.add([SharedFile.from_mb90(base, rec, row["raw_text"])])
        summary = f"{rec['record_type']} {rec.get('end_result', '')} {rec.get('end_result_unit', '')}"
    else:
        rec = parse_ph(text)
        row = {**base, **rec, "raw_text": text.replace("\n", " | ")}
        WRITER.append(monthly_csv(out_dir, name), [row])
        SHARED.add([SharedFile.from_ph(base, rec, row["raw_text"])])
        summary = f"pH {rec['pH'] or '-'}  {rec['temperature']}{rec['temp_unit']}  mV {rec['mV'] or '-'}"
    log(f"{name}: saved {summary.strip()}  (lot '{ctx['lot']}')")


# --------------------------------------------------------------------------
# Colorimeter folder watcher
# --------------------------------------------------------------------------
class FolderWatcher(threading.Thread):
    EXTENSIONS = {".csv", ".xlsx", ".xlsm"}

    def __init__(self, cfg, out_dir):
        super().__init__(daemon=True)
        self.cfg = cfg
        self.name_ = cfg["name"]
        self.out_dir = out_dir
        self.folder = expand(cfg["watch_folder"])
        self.done_dir = self.folder / "imported"
        self.poll = float(cfg.get("poll_seconds", 10))
        self.skip_dupes = bool(cfg.get("skip_duplicates", False))
        self.seen_path = out_dir / self.name_ / ".imported_row_ids.txt"
        self.seen = set()
        if self.skip_dupes and self.seen_path.exists():
            self.seen = set(self.seen_path.read_text(encoding="utf-8").split())
        set_status(self.name_, f"Watching {self.folder.name}", True, type=cfg.get("type"),
                   label=cfg.get("label") or DEVICE_LABELS.get(cfg.get("type"), self.name_))

    @staticmethod
    def row_id(header, r):
        import hashlib
        parts = [f"{k}={str(v).strip()}" for k, v in zip(header, r)]
        return hashlib.sha1("|".join(parts).encode("utf-8")).hexdigest()[:20]

    def run(self):
        self.folder.mkdir(parents=True, exist_ok=True)
        self.done_dir.mkdir(exist_ok=True)
        log(f"{self.name_}: watching folder {self.folder}")
        while not STOP.is_set():
            for f in sorted(self.folder.iterdir()):
                if f.is_file() and f.suffix.lower() in self.EXTENSIONS and not f.name.startswith("~$"):
                    if time.time() - f.stat().st_mtime < 3:
                        continue  # still being written
                    try:
                        self.import_file(f)
                    except PermissionError:
                        continue  # still open in the Linshang software / Excel
                    except Exception as e:
                        log(f"{self.name_}: could not read {f.name}: {e}")
                        self.move(f, self.folder / "failed")
            STOP.wait(self.poll)

    def read_rows(self, f):
        if f.suffix.lower() == ".csv":
            with open(f, newline="", encoding="utf-8-sig", errors="replace") as fh:
                table = [r for r in csv.reader(fh)]
            return [("csv", table)]
        if openpyxl is None:
            raise RuntimeError("openpyxl not installed (pip install openpyxl)")
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        sheets = []
        for ws in wb.worksheets:
            sheets.append((ws.title, [["" if c is None else str(c) for c in row]
                                      for row in ws.iter_rows(values_only=True)]))
        wb.close()
        return sheets

    def import_file(self, f):
        ctx = context_snapshot()
        rows_out, new_ids, dupes = [], [], 0
        for sheet, table in self.read_rows(f):
            table = [r for r in table if any(str(c).strip() for c in r)]
            if not table:
                continue
            # header = first row with at least 2 non-empty cells
            h_idx = next((i for i, r in enumerate(table)
                          if sum(1 for c in r if str(c).strip()) >= 2), 0)
            top = [str(c).strip() for c in table[h_idx]]
            data_start = h_idx + 1
            # Two-row headers (Linshang: "Standard" / "Sample" above L* a* b*): merge them,
            # giving "Standard L*", "Sample L*" and so on.
            nxt = [str(c).strip() for c in table[h_idx + 1]] if len(table) > h_idx + 1 else []
            if nxt and sum(1 for c in nxt if c) >= 2 and not any(_num(c) is not None for c in nxt if c):
                group, merged = "", []
                for i in range(max(len(top), len(nxt))):
                    t = top[i] if i < len(top) else ""
                    b = nxt[i] if i < len(nxt) else ""
                    if t:
                        group = t if b else ""
                    merged.append(f"{group} {b}".strip() if b else t)
                top, data_start = merged, h_idx + 2
            header, seen = [], {}
            for i, c in enumerate(top):
                c = c or f"col{i + 1}"
                seen[c] = seen.get(c, 0) + 1
                header.append(c if seen[c] == 1 else f"{c}_{seen[c]}")
            for n, r in enumerate(table[data_start:], start=1):
                if self.skip_dupes:
                    rid = self.row_id(header, r)
                    if rid in self.seen or rid in new_ids:
                        dupes += 1
                        continue
                    new_ids.append(rid)
                row = {"logged_at": now_iso(), "instrument": self.name_,
                       "lot": ctx["lot"], "operator": ctx["operator"], "note": ctx["note"],
                       "source_file": f.name, "sheet": sheet, "row": n}
                for k, v in zip(header, r):
                    row[k] = v
                rows_out.append(row)
        WRITER.append(monthly_csv(self.out_dir, self.name_), rows_out)
        if self.cfg.get("type") == "export_folder":
            SHARED.add([SHARED.from_aqualab_export(r) for r in rows_out])
        else:
            SHARED.add([SharedFile.from_colorimeter(r) for r in rows_out])
        if new_ids:
            self.seen.update(new_ids)
            self.seen_path.parent.mkdir(parents=True, exist_ok=True)
            with open(self.seen_path, "a", encoding="utf-8") as fh:
                fh.write("\n".join(new_ids) + "\n")
        dest = self.move(f, self.done_dir)
        extra = f", skipped {dupes} already-imported row(s)" if dupes else ""
        log(f"{self.name_}: imported {len(rows_out)} row(s) from {f.name}{extra} -> {dest.parent.name}/")

    @staticmethod
    def move(f, target_dir):
        target_dir.mkdir(exist_ok=True)
        dest = target_dir / f"{dt.datetime.now():%Y%m%d-%H%M%S}_{f.name}"
        shutil.move(str(f), dest)
        return dest



# --------------------------------------------------------------------------
# SKALA Control fetcher (AQUALAB 3 readings, read-only API calls)
# --------------------------------------------------------------------------
SKALA_SYNC_NOW = threading.Event()
SKALA_MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}


def parse_skala_datetime(text):
    """'30/Sept/2026 11:41' -> '2026-09-30T11:41:00' (returns '' if unparseable)."""
    m = re.match(r"\s*(\d{1,2})/([A-Za-z]+)/(\d{4})(?:\s+(\d{1,2}):(\d{2}))?", str(text or ""))
    if not m:
        return ""
    month = SKALA_MONTHS.get(m.group(2)[:3].lower())
    if not month:
        return ""
    try:
        return dt.datetime(int(m.group(3)), month, int(m.group(1)),
                           int(m.group(4) or 0), int(m.group(5) or 0)).isoformat()
    except ValueError:
        return ""


class SkalaAuthError(Exception):
    pass


class SkalaClient:
    """Makes the same read-only requests the SKALA Control web page makes."""

    def __init__(self, api_base, token, sbs_serial, timeout=30):
        self.api_base = api_base.rstrip("/")
        self.token = token
        self.sbs_serial = sbs_serial
        self.timeout = timeout

    def call(self, endpoint, **body):
        import urllib.request
        import urllib.error
        body.setdefault("sbs_serial", self.sbs_serial)
        req = urllib.request.Request(
            f"{self.api_base}/{endpoint}",
            data=json.dumps(body).encode("utf-8"),
            headers={"token": self.token, "Accept": "application/json",
                     "Content-Type": "application/json"},
            method="POST",
        )
        try:
            with urllib.request.urlopen(req, timeout=self.timeout) as r:
                raw = r.read().decode("utf-8", errors="replace")
        except urllib.error.HTTPError as e:
            if e.code in (401, 403):
                raise SkalaAuthError(f"SKALA refused the token (HTTP {e.code})")
            raise
        try:
            outer = json.loads(raw)
        except json.JSONDecodeError:
            raise RuntimeError(f"{endpoint}: unexpected reply {raw[:200]!r}")
        payload = outer.get("payload") if isinstance(outer, dict) else outer
        if isinstance(payload, str):
            try:
                payload = json.loads(payload)
            except json.JSONDecodeError:
                pass
        if isinstance(payload, str) and re.search(r"token|auth|login|expired|unauthori", payload, re.I):
            raise SkalaAuthError(f"SKALA says: {payload[:150]}")
        if isinstance(outer, dict) and re.search(r"token|unauthori|expired",
                                                 str(outer.get("message", outer.get("error", ""))), re.I):
            raise SkalaAuthError(f"SKALA says: {outer.get('message') or outer.get('error')}")
        return payload

    def paged(self, endpoint, list_key, total_key, page_size=50, max_pages=50, **body):
        page = 1
        got = 0
        while page <= max_pages:
            data = self.call(endpoint, page_num=page, page_size=page_size, **body) or {}
            items = (data.get(list_key) or []) if isinstance(data, dict) else []
            for it in items:
                yield it
            got += len(items)
            total = data.get(total_key) if isinstance(data, dict) else None
            if not items or len(items) < page_size or (isinstance(total, int) and got >= total):
                return
            page += 1


class SkalaFetcher(threading.Thread):
    """Polls SKALA Control and saves new AQUALAB 3 readings. Only reads data."""

    def __init__(self, cfg, out_dir):
        super().__init__(daemon=True)
        self.cfg = cfg
        self.name_ = cfg["name"]
        self.out_dir = out_dir
        self.api_base = cfg.get("api_base", "https://lambda-prod.skalacontrol.com")
        self.token_file = expand(cfg.get("token_file", str(SCRIPT_DIR / "skala_token.txt")))
        self.sbs_serial = cfg.get("sbs_serial", "")
        self.poll = max(1.0, float(cfg.get("poll_minutes", 5))) * 60
        self.lookback = float(cfg.get("lookback_days", 3))
        self.page_size = int(cfg.get("page_size", 50))
        self.full_rescan = float(cfg.get("full_rescan_minutes", 60)) * 60
        self.state_path = out_dir / self.name_ / ".skala_state.json"
        self.seen_path = out_dir / self.name_ / ".skala_reading_ids.txt"
        self.seen = set()
        if self.seen_path.exists():
            self.seen = set(self.seen_path.read_text(encoding="utf-8").split())
        self.state = {"samples": {}, "last_full_scan": 0}
        if self.state_path.exists():
            try:
                self.state = json.loads(self.state_path.read_text(encoding="utf-8"))
            except Exception:
                pass
        self.auth_warned = False
        set_status(self.name_, "SKALA · starting", False, type="skala_api",
                   label=cfg.get("label") or DEVICE_LABELS["skala_api"])

    def read_token(self):
        try:
            return self.token_file.read_text(encoding="utf-8").strip().strip('"').strip("'")
        except FileNotFoundError:
            return ""

    def run(self):
        log(f"{self.name_}: SKALA fetcher on, checking every {self.poll / 60:g} min "
            f"(type 'sync' to check now)")
        while not STOP.is_set():
            try:
                n = self.poll_once()
                set_status(self.name_, f"SKALA · last check {dt.datetime.now():%H:%M}", True)
                if n:
                    log(f"{self.name_}: saved {n} new reading(s) from SKALA")
                self.auth_warned = False
            except SkalaAuthError as e:
                set_status(self.name_, "SKALA · token expired or invalid", False)
                if not self.auth_warned:
                    log(f"{self.name_}: {e}. Paste a fresh token into {self.token_file} "
                        f"(see guide); will keep retrying.")
                    self.auth_warned = True
            except Exception as e:  # offline, timeout, server hiccup
                set_status(self.name_, "SKALA · unreachable (no internet?)", False)
                if "urlopen" in str(e) or "timed out" in str(e) or isinstance(e, OSError):
                    pass  # quietly wait while offline
                else:
                    log(f"{self.name_}: SKALA check failed: {e}")
            SKALA_SYNC_NOW.wait(self.poll)
            SKALA_SYNC_NOW.clear()

    def poll_once(self, verbose=False):
        token = self.read_token()
        if not token:
            raise SkalaAuthError(f"No token in {self.token_file}")
        if not self.sbs_serial:
            raise RuntimeError("sbs_serial is empty in the config")
        client = SkalaClient(self.api_base, token, self.sbs_serial)
        cutoff = dt.datetime.now() - dt.timedelta(days=self.lookback)
        full = time.time() - self.state.get("last_full_scan", 0) > self.full_rescan
        new_rows, new_ids = [], []

        for batch in self.recent_batches(client, cutoff):
            bdate = str(batch.get("date") or "")
            if verbose:
                print(f"  batch {batch.get('id')} '{batch.get('name')}' ({batch.get('product')}, {bdate})")
            for sample in client.paged("get_samples_by_batch", "samples", "total_samples",
                                       self.page_size, batch_id=batch["id"], sample_order="latest"):
                sid = str(sample.get("id"))
                marker = f"{sample.get('last_reading_datetime')}|{sample.get('number_completed_required_tests')}"
                if not full and self.state["samples"].get(sid) == marker:
                    continue
                sample_fields = "; ".join(
                    f"{f.get('name') or 'field'}={f.get('value')}" for f in sample.get("fields") or [])
                tests = (client.call("get_tests_by_sample", sample_id=sample["id"]) or {})
                for test in tests.get("required_tests") or []:
                    for rd in client.paged("get_readings_by_test", "readings", "total_readings",
                                           self.page_size, sample_id=sample["id"],
                                           required_test_id=test["id"]):
                        rid = str(rd.get("id"))
                        if rid in self.seen or rid in new_ids:
                            continue
                        new_ids.append(rid)
                        new_rows.append(self.make_row(rd, batch, sample, test, sample_fields))
                self.state["samples"][sid] = marker

        if new_rows:
            new_rows.sort(key=lambda r: r["reading_datetime"] or "")
            WRITER.append(monthly_csv(self.out_dir, self.name_), new_rows)
            SHARED.add([SHARED.from_aqualab(r) for r in new_rows])
            self.seen.update(new_ids)
            self.seen_path.parent.mkdir(parents=True, exist_ok=True)
            with open(self.seen_path, "a", encoding="utf-8") as fh:
                fh.write("\n".join(new_ids) + "\n")
        if full:
            self.state["last_full_scan"] = time.time()
        self.state_path.parent.mkdir(parents=True, exist_ok=True)
        self.state_path.write_text(json.dumps(self.state), encoding="utf-8")
        return len(new_rows)

    def recent_batches(self, client, cutoff, max_pages=40):
        """Batches dated or changed within the lookback window. Stops after a
        whole page of batches that are all older than the window."""
        cut = cutoff.isoformat()
        for page in range(1, max_pages + 1):
            data = client.call("get_batches_ionic", page_num=page, page_size=self.page_size,
                               filter_text="") or {}
            batches = (data.get("batches") or []) if isinstance(data, dict) else []
            any_recent = False
            for b in batches:
                bdate = str(b.get("date") or "")
                changed = parse_skala_datetime(b.get("last_change"))
                newest = max(filter(None, [changed, bdate + "T23:59:59" if bdate else ""]), default="")
                if not newest or newest >= cut:
                    any_recent = True
                    yield b
            total = data.get("total_batches") if isinstance(data, dict) else None
            if not batches or len(batches) < self.page_size or not any_recent:
                return
            if isinstance(total, int) and page * self.page_size >= total:
                return

    def make_row(self, rd, batch, sample, test, sample_fields):
        ctx = context_snapshot()
        is_aw = rd.get("unit_class") == "WaterActivityUnit"
        return {
            "logged_at": now_iso(),
            "instrument": self.name_,
            "source": "SKALA Control",
            "reading_id": rd.get("id"),
            "reading_datetime": parse_skala_datetime(rd.get("datetime")),
            "reading_datetime_skala": rd.get("datetime"),
            "product": rd.get("product_name") or batch.get("product"),
            "batch_name": batch.get("name"),
            "batch_date": batch.get("date"),
            "batch_id": rd.get("batch_id") or batch.get("id"),
            "sample_id": rd.get("sample_id") or sample.get("id"),
            "sample_fields": sample_fields,
            "test": rd.get("translated_reading_type") or rd.get("reading_type") or test.get("unit_name"),
            "water_activity": rd.get("primary_reading") if is_aw else "",
            "temperature_C": rd.get("secondary_reading") if is_aw else "",
            "primary_reading": rd.get("primary_reading"),
            "secondary_reading": rd.get("secondary_reading"),
            "tertiary_reading": rd.get("tertiary_reading"),
            "dual_block_difference": rd.get("dual_block_difference"),
            "in_spec": rd.get("in_spec"),
            "in_haccp_spec": rd.get("in_haccp_spec"),
            "spec_min": test.get("min_value"),
            "spec_max": test.get("max_value"),
            "status": rd.get("status"),
            "void": rd.get("void"),
            "is_fast": rd.get("is_fast"),
            "duration": rd.get("duration"),
            "instrument_sn": rd.get("sn"),
            "instrument_type": rd.get("devType"),
            "instrument_verified_at": rd.get("instrument_verified_at"),
            "skala_user_id": rd.get("user_id"),
            "supervisor_approved": rd.get("supervisor_approved"),
            "tag": rd.get("tag"),
            "comment": rd.get("comment"),
            "lot": ctx["lot"], "operator": ctx["operator"], "note": ctx["note"],
            "raw_json": json.dumps(rd, separators=(",", ":"), default=str),
        }

# --------------------------------------------------------------------------
# Excel workbook (always on) and daily history snapshots
# --------------------------------------------------------------------------
def _cell(v):
    """Numbers stay numbers in Excel so people can sort and average them."""
    if isinstance(v, str):
        t = v.strip()
        if re.fullmatch(r"-?\d+(\.\d+)?", t) and not re.fullmatch(r"0\d+", t):
            try:
                return float(t) if "." in t else int(t)
            except ValueError:
                return v
    return v


class ExcelBook:
    """Keeps data/Excel/Lab_Results_<month>.xlsx up to date from the saved records.
    The workbook is a view: it's rebuilt from the CSV records, which are never edited."""

    def __init__(self):
        self.dirty = True
        self.months = set()
        self.last_build = 0.0
        self.updated_at = ""
        self.locked = False
        self.folder = None
        self.out_dir = None
        self.lock = threading.Lock()

    def configure(self, cfg, out_dir):
        self.out_dir = out_dir
        self.folder = out_dir / "Excel"
        self.folder.mkdir(parents=True, exist_ok=True)
        self.months = {r.get("_month") for r in read_all_results() if r.get("_month")}
        self.months.add(f"{dt.date.today():%Y-%m}")
        self.dirty = True

    def path(self, month=None):
        return self.folder / f"Lab_Results_{month or f'{dt.date.today():%Y-%m}'}.xlsx"

    def mark(self, month=None):
        self.months.add(month or f"{dt.date.today():%Y-%m}")
        self.dirty = True

    def build(self, month, dest):
        if openpyxl is None:
            raise RuntimeError("openpyxl is not installed")
        from openpyxl.styles import Font, Alignment, PatternFill
        bold = Font(name="Arial", bold=True)
        head_fill = PatternFill("solid", fgColor="E7EBF0")
        rows = [r for r in read_all_results() if r.get("_month") == month]
        wb = openpyxl.Workbook()

        def sheet(title, fields, data, first=False):
            ws = wb.active if first else wb.create_sheet()
            ws.title = title[:31]
            ws.append(fields)
            for d in data:
                ws.append([_cell(d.get(f, "")) for f in fields])
            for c in ws[1]:
                c.font = bold
                c.fill = head_fill
                c.alignment = Alignment(vertical="center", wrap_text=True)
            ws.freeze_panes = "A2"
            if ws.max_row > 1:
                ws.auto_filter.ref = ws.dimensions
            for i, f in enumerate(fields, 1):
                width = max([len(str(f))] + [len(str(d.get(f, ""))) for d in data[:200]]) + 2
                ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = min(max(width, 9), 40)
            # Protect against accidental typing (no password): sorting and filters still work.
            ws.protection.sheet = True
            ws.protection.autoFilter = False
            ws.protection.sort = False
            ws.protection.formatColumns = False
            return ws

        fields, summary = summarize_rows(rows)
        sheet("Lot Summary", fields, summary, first=True)
        extra = ["lot_assigned_later", "assigned_by", "assigned_at"]
        sheet("All Results", SHARED.columns + [e for e in extra if e not in SHARED.columns], rows)
        # each instrument's full detail
        for sub in sorted(p for p in self.out_dir.iterdir() if p.is_dir()):
            for f in (sub / f"{sub.name}_{month}.csv", sub / f"{sub.name}_curves_{month}.csv"):
                if f.exists():
                    with open(f, newline="", encoding="utf-8-sig") as fh:
                        rd = csv.DictReader(fh)
                        data = list(rd)
                        flds = [c for c in (rd.fieldnames or []) if c != "raw_json"]
                    label = (DEVICES.get(sub.name) or {}).get("label") or sub.name
                    sheet(label + (" curves" if "_curves_" in f.name else ""), flds, data)
        af = SHARED.folder / "Lot_Assignments.csv"
        if af.exists():
            with open(af, newline="", encoding="utf-8-sig") as fh:
                rd = csv.DictReader(fh)
                data = [a for a in rd if (a.get("assigned_at") or "")[:7] == month]
                flds = rd.fieldnames or []
            sheet("Lot Assignments", flds, data)
        hist = [h for h in (load_upload_state().get("history") or []) if (h.get("when") or "")[:7] == month]
        sheet("Uploads", ["when", "scope", "results", "files", "dest", "operator"], hist)
        ws = wb.create_sheet("About")
        for line in [
            "Lab Logger results for " + month,
            "Updated " + now_iso().replace("T", " "),
            "",
            "This workbook is rebuilt automatically from the permanent records, so anything typed here is replaced.",
            "To keep your own notes or charts, save a copy (File > Save As) under a different name.",
            "Permanent records (only ever added to): " + str(self.out_dir),
            "Daily backups: " + str(self.out_dir / "History"),
        ]:
            ws.append([line])
        ws["A1"].font = Font(name="Arial", bold=True, size=13)
        ws.column_dimensions["A"].width = 110
        ws.protection.sheet = True
        dest = Path(dest)
        dest.parent.mkdir(parents=True, exist_ok=True)
        tmp = dest.with_name(dest.stem + ".tmp.xlsx")
        wb.save(tmp)
        os.replace(tmp, dest)

    def rebuild_if_needed(self, min_gap=10):
        if not self.dirty or time.time() - self.last_build < min_gap or WRITER.pending:
            return
        with self.lock:
            months, self.dirty = set(self.months), False
            self.months = {f"{dt.date.today():%Y-%m}"}
            blocked = []
            for m in sorted(months):
                try:
                    self.build(m, self.path(m))
                except PermissionError:
                    blocked.append(m)
                except Exception as e:
                    log(f"Excel workbook {m} could not be updated: {e}")
            self.last_build = time.time()
            if blocked:
                self.locked = True
                self.months |= set(blocked)
                self.dirty = True
            else:
                self.locked = False
                self.updated_at = now_iso()

    def open_copy(self):
        """Open a fresh copy in Excel, so the live workbook never gets locked."""
        import tempfile
        month = f"{dt.date.today():%Y-%m}"
        dest = Path(tempfile.gettempdir()) / f"Lab_Results_{month}_{dt.datetime.now():%H%M%S}.xlsx"
        self.build(month, dest)
        if not hasattr(os, "startfile"):
            raise RuntimeError(f"Opening Excel only works on Windows. A copy was saved to {dest}")
        try:
            os.startfile(str(dest))
        except OSError:
            raise RuntimeError("This laptop doesn't have a program that opens Excel files. "
                               "Your results are still being saved. To view them here, install the free "
                               f"LibreOffice, or ask {DASH['support']} to open the file on another computer.")
        return str(dest)


EXCEL = ExcelBook()


class History:
    """Once a day, zip the whole data folder into data/History. Kept forever unless keep_days is set."""

    def __init__(self):
        self.enabled = True
        self.keep_days = 0
        self.folder = None
        self.out_dir = None
        self.last = ""
        self.next_check = 0.0

    def configure(self, cfg, out_dir):
        hc = cfg.get("history", {})
        self.enabled = hc.get("enabled", True)
        self.keep_days = int(hc.get("keep_days", 0) or 0)
        self.out_dir = out_dir
        self.folder = out_dir / "History"
        self.folder.mkdir(parents=True, exist_ok=True)
        snaps = sorted(self.folder.glob("Lab_Data_*.zip"))
        self.last = snaps[-1].name if snaps else ""

    def check(self, force=False):
        """Keep today's backup at most an hour behind. Earlier days' backups are never touched."""
        if not self.enabled or (not force and time.time() < self.next_check):
            return
        self.next_check = time.time() + 600
        target = self.folder / f"Lab_Data_{dt.date.today():%Y-%m-%d}.zip"
        if target.exists():
            made = target.stat().st_mtime
            if not force and time.time() - made < 3600:
                return
            changed = any(f.is_file() and f.stat().st_mtime > made for f in self.out_dir.rglob("*")
                          if f.relative_to(self.out_dir).parts[0] != "History")
            if not changed:
                return
        import zipfile
        tmp = target.with_suffix(".tmp")
        try:
            with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
                for f in self.out_dir.rglob("*"):
                    rel = f.relative_to(self.out_dir)
                    if f.is_file() and rel.parts[0] not in ("History",) and not f.name.endswith(".tmp"):
                        try:
                            z.write(f, rel)
                        except (PermissionError, OSError):
                            pass
            os.replace(tmp, target)
            self.last = target.name
            if not force:
                log(f"Daily backup saved: {target.name}")
        except Exception as e:
            log(f"Daily backup failed: {e}")
            return
        if self.keep_days > 0:
            cutoff = f"Lab_Data_{dt.date.today() - dt.timedelta(days=self.keep_days):%Y-%m-%d}.zip"
            for old in self.folder.glob("Lab_Data_*.zip"):
                if old.name < cutoff:
                    old.unlink()


HISTORY = History()


def housekeeping():
    """Runs every few seconds: Lot Summary, Excel workbook, daily backup."""
    if not WRITER.pending:
        SHARED.rebuild_summary()
    EXCEL.rebuild_if_needed()
    HISTORY.check()


# --------------------------------------------------------------------------
# Lot rules: lot/operator expire, results without a lot wait for one
# --------------------------------------------------------------------------
NO_LOT_NEEDED = re.compile(r"calibration|statistics|printout", re.I)


def configure_lot_rules(cfg):
    LOT_RULES.update(cfg.get("lot_rules", {}))


def context_expires_at():
    with CONTEXT_LOCK:
        set_at = CONTEXT.get("set_at")
        active = CONTEXT.get("lot") or CONTEXT.get("operator")
    if not set_at or not active:
        return ""
    t = dt.datetime.fromisoformat(set_at)
    ends = []
    if LOT_RULES.get("expire_minutes"):
        ends.append(t + dt.timedelta(minutes=float(LOT_RULES["expire_minutes"])))
    if LOT_RULES.get("clear_at_midnight", True):
        ends.append(dt.datetime.combine(t.date() + dt.timedelta(days=1), dt.time()))
    return min(ends).isoformat(timespec="seconds") if ends else ""


def clear_context(reason):
    with CONTEXT_LOCK:
        had = CONTEXT.get("lot") or CONTEXT.get("operator")
        CONTEXT.update(lot="", operator="", note="", set_at="")
    if had:
        log(f"Lot and operator cleared ({reason}). They must be entered again before the next test.")


def context_watch():
    while not STOP.wait(15):
        exp = context_expires_at()
        if exp and now_iso() >= exp:
            clear_context("time limit reached")


def check_context(lot, operator):
    """Return a plain-words problem with a lot/operator, or ''."""
    if not lot:
        return "Enter the lot number."
    if not operator:
        return "Enter your name."
    pat = LOT_RULES.get("lot_pattern")
    if pat and not re.fullmatch(pat, lot):
        ex = LOT_RULES.get("lot_example")
        return f"'{lot}' doesn't look like a lot number" + (f" (they look like {ex})." if ex else ".")
    ops = LOT_RULES.get("operators") or []
    if ops and operator not in ops:
        return "Choose your name from the list."
    return ""


def needs_lot(row):
    if NO_LOT_NEEDED.search(row.get("measurement") or ""):
        return False
    if not (row.get("lot") or "").strip():
        return True
    if row.get("skala_reading_id"):      # AQUALAB: SKALA already records who ran it
        return False
    return not (row.get("operator") or "").strip()


def results_needing_lot():
    out = []
    for r in read_all_results():
        if needs_lot(r):
            out.append({"record_id": r.get("record_id", ""), "time": r.get("reading_datetime") or r.get("logged_at", ""),
                        "instrument": r.get("instrument", ""), "measurement": r.get("measurement", ""),
                        "value": r.get("value", ""), "unit": r.get("unit", ""),
                        "lot": r.get("lot", ""), "operator": r.get("operator", "")})
    return out


def assign_lot(record_ids, lot, operator):
    """Give a lot/operator to results that arrived without one. Nothing already saved is
    edited: each assignment is added to Lot_Assignments.csv and applied when files are read."""
    problem = check_context(lot, operator)
    if problem:
        raise ValueError(problem)
    ids = set(record_ids)
    changed = []
    for r in read_all_results():
        if r.get("record_id") in ids and needs_lot(r):
            new_lot = (r.get("lot") or "").strip() or lot
            new_op = (r.get("operator") or "").strip() or operator
            changed.append({"assigned_at": now_iso(), "record_id": r["record_id"],
                            "instrument": r.get("instrument", ""), "measurement": r.get("measurement", ""),
                            "value": r.get("value", ""), "lot_before": r.get("lot", ""),
                            "operator_before": r.get("operator", ""), "lot": new_lot,
                            "operator": new_op, "assigned_by": operator})
    if changed:
        WRITER.append(SHARED.folder / "Lot_Assignments.csv", changed)
        WRITER.flush()
        SHARED.dirty = True
        for r in read_all_results():
            if r.get("record_id") in ids:
                EXCEL.mark(r.get("_month"))
        EXCEL.mark()
        with DEV_LOCK:
            byid = {c["record_id"]: c for c in changed}
            for item in RECENT:
                c = byid.get(item.get("record_id"))
                if c:
                    item.update(lot=c["lot"], operator=c["operator"])
        log(f"Assigned {len(changed)} result(s) to lot '{lot}' by {operator}")
    return len(changed)


# --------------------------------------------------------------------------
# Dashboard: a local web page at http://localhost:8080 (this laptop only)
# --------------------------------------------------------------------------
UPLOAD_CFG = {"folder": ""}
DASH = {"out_dir": None, "station": "", "port": 8080, "support": "your supervisor"}


def summarize_rows(rows):
    """One summary row per lot (same rules as Lot_Summary.csv)."""
    lots = {}
    for r in rows:
        lot = (r.get("lot") or "").strip() or "(no lot)"
        d = lots.setdefault(lot, {"times": [], "vals": {}, "n": 0})
        d["n"] += 1
        tested = r.get("reading_datetime") or r.get("logged_at")
        if tested:
            d["times"].append(tested)
        for title, meas, field in SUMMARY_RESULTS:
            if r.get("measurement") == meas:
                v = _num(r.get(field))
                if v is not None:
                    d["vals"].setdefault(title, []).append(v)
    out = []
    for lot, d in lots.items():
        t = sorted(d["times"])
        first, last = (t[0], t[-1]) if t else ("", "")
        row = {"Lot": lot, "First date": first[:10], "First time": first[11:16],
               "Last date": last[:10], "Last time": last[11:16]}
        for title, _, _ in SUMMARY_RESULTS:
            vals = d["vals"].get(title)
            row[title] = round(sum(vals) / len(vals), 4) if vals else ""
        row["Results logged"] = d["n"]
        out.append(row)
    out.sort(key=lambda r: (r["First date"], r["First time"]), reverse=True)
    fields = (["Lot", "First date", "First time", "Last date", "Last time"]
              + [t for t, _, _ in SUMMARY_RESULTS] + ["Results logged"])
    return fields, out


_ALL_CACHE = {"key": None, "rows": []}


def load_assignments():
    """record_id -> latest assignment (from Lot_Assignments.csv, which is only ever added to)."""
    f = SHARED.folder / "Lot_Assignments.csv" if SHARED.folder else None
    out = {}
    if f and f.exists():
        try:
            with open(f, newline="", encoding="utf-8-sig") as fh:
                for a in csv.DictReader(fh):
                    out[a.get("record_id", "")] = a
        except Exception:
            pass
    return out


def read_all_results():
    """All shared results with later lot assignments applied. Re-read only when files change.
    The saved files themselves are never edited."""
    if not SHARED.enabled or not SHARED.folder:
        return []
    files = sorted(SHARED.folder.glob("All_Results_*.csv"))
    extra = [SHARED.folder / "Lot_Assignments.csv"]
    key = tuple((f.name, f.stat().st_mtime_ns, f.stat().st_size) for f in files + extra if f.exists())
    if key == _ALL_CACHE["key"]:
        return _ALL_CACHE["rows"]
    rows = []
    for f in files:
        month = f.stem.replace("All_Results_", "")
        try:
            with open(f, newline="", encoding="utf-8-sig") as fh:
                for r in csv.DictReader(fh):
                    r["_month"] = month
                    rows.append(r)
        except Exception:
            return _ALL_CACHE["rows"]
    assigned = load_assignments()
    for r in rows:
        a = assigned.get(r.get("record_id"))
        if a:
            if not (r.get("lot") or "").strip():
                r["lot"] = a.get("lot", "")
            if not (r.get("operator") or "").strip():
                r["operator"] = a.get("operator", "")
            r["lot_assigned_later"] = "yes"
            r["assigned_by"] = a.get("assigned_by", "")
            r["assigned_at"] = a.get("assigned_at", "")
    _ALL_CACHE.update(key=key, rows=rows)
    return rows


def upload_state_path():
    return DASH["out_dir"] / ".upload_state.json"


def load_upload_state():
    try:
        return json.loads(upload_state_path().read_text(encoding="utf-8"))
    except Exception:
        return {"last_upload": "", "history": []}


def pending_count():
    last = load_upload_state().get("last_upload", "")
    return sum(1 for r in read_all_results() if (r.get("logged_at") or "") > last)


def select_rows(scope, lot="", date_from="", date_to=""):
    rows = read_all_results()
    if scope == "lot":
        return [r for r in rows if (r.get("lot") or "").strip() == lot.strip()]
    if scope == "range":
        def day(r):
            return (r.get("reading_datetime") or r.get("logged_at") or "")[:10]
        return [r for r in rows if (not date_from or day(r) >= date_from) and (not date_to or day(r) <= date_to)]
    last = load_upload_state().get("last_upload", "")
    return [r for r in rows if (r.get("logged_at") or "") > last]


def build_upload_workbook(rows, path):
    if openpyxl is None:
        raise RuntimeError("openpyxl is not installed (pip install openpyxl)")
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    fields, summary = summarize_rows(rows)
    ws = wb.active
    ws.title = "Lot Summary"
    ws.append(fields)
    for r in summary:
        ws.append([r.get(f, "") for f in fields])
    ws2 = wb.create_sheet("All Results")
    cols = SHARED.columns
    ws2.append(cols)
    for r in rows:
        ws2.append([r.get(c, "") for c in cols])
    for sheet in (ws, ws2):
        for c in sheet[1]:
            c.font = Font(name="Arial", bold=True)
        sheet.freeze_panes = "A2"
    wb.save(path)


def do_upload(req):
    """req: scope (new|lot|range), lot, from, to, include [shared, detail, raw], dest (folder|download).
    Returns (info dict, file path to send or None)."""
    import shutil as _sh
    import tempfile
    import zipfile
    scope = req.get("scope", "new")
    rows = select_rows(scope, req.get("lot", ""), req.get("from", ""), req.get("to", ""))
    missing = sum(1 for r in rows if needs_lot(r))
    if missing and LOT_RULES.get("block_upload_without_lot", True):
        raise RuntimeError(f"{missing} result{'s' if missing != 1 else ''} still need{'s' if missing == 1 else ''} a lot number. "
                           "Assign them on the main page first, then upload.")
    include = set(req.get("include") or ["shared"])
    stamp = dt.datetime.now().strftime("%Y-%m-%d_%H%M")
    label = {"new": "new", "lot": "lot-" + re.sub(r"[^\w-]+", "_", req.get("lot", "")),
             "range": f"{req.get('from', '')}_to_{req.get('to', '')}"}.get(scope, "new")
    out_dir = DASH["out_dir"]
    tmp = Path(tempfile.mkdtemp(prefix="lablogger_upload_"))
    files = []
    if "shared" in include:
        x = tmp / f"Lab_Results_{stamp}_{label}.xlsx"
        build_upload_workbook(rows, x)
        files.append((x, x.name))
    for sub in [p for p in out_dir.iterdir() if p.is_dir() and p.name not in ("Shared", "Excel", "History")
                and not p.name.endswith("Exports")] if include & {"detail", "raw"} else []:
        for f in sub.glob("*.csv"):
            if "detail" in include:
                files.append((f, f"{sub.name}/{f.name}"))
        if "raw" in include and (sub / "raw").is_dir():
            for f in (sub / "raw").glob("*.txt"):
                files.append((f, f"{sub.name}/raw/{f.name}"))
    if req.get("dest") == "folder":
        dest = UPLOAD_CFG.get("folder")
        if not dest:
            raise RuntimeError(f"Uploading to the cloud folder isn't set up yet. Use Download for now, or ask {DASH['support']}.")
        dest = expand(dest)
        target = dest / f"Upload_{stamp}_{label}"
        for src, rel in files:
            (target / rel).parent.mkdir(parents=True, exist_ok=True)
            _sh.copy2(src, target / rel)
        sent = None
        where = str(target)
    else:
        if len(files) == 1 and files[0][1].endswith(".xlsx"):
            sent = files[0][0]
        else:
            sent = tmp / f"Lab_Data_{stamp}_{label}.zip"
            with zipfile.ZipFile(sent, "w", zipfile.ZIP_DEFLATED) as z:
                for src, rel in files:
                    z.write(src, rel)
        where = "download"
    st = load_upload_state()
    if scope == "new":
        st["last_upload"] = now_iso()
    ctx = context_snapshot()
    st.setdefault("history", []).insert(0, {
        "when": now_iso(), "scope": label, "results": len(rows), "files": len(files),
        "dest": "cloud folder" if req.get("dest") == "folder" else "download",
        "operator": ctx.get("operator", "")})
    st["history"] = st["history"][:30]
    upload_state_path().write_text(json.dumps(st, indent=1), encoding="utf-8")
    EXCEL.mark()
    log(f"Upload: {len(rows)} result(s), {len(files)} file(s) -> {where}")
    return {"ok": True, "results": len(rows), "files": len(files), "where": where}, sent


def dashboard_state():
    with DEV_LOCK:
        devices = []
        for name, d in DEVICES.items():
            item = {k: v for k, v in d.items()}
            lv = LIVE.get(name)
            item["live_active"] = bool(lv and lv.get("test") and lv["test"].get("active"))
            devices.append(item)
        recent = list(RECENT)[:25]
        ctx = dict(CONTEXT)
    needs = results_needing_lot()
    return {"station": DASH["station"], "support": DASH["support"], "context": ctx,
            "expires_at": context_expires_at(), "needs_lot": needs[-50:], "needs_lot_count": len(needs),
            "operators": LOT_RULES.get("operators") or [], "lot_example": LOT_RULES.get("lot_example", ""),
            "block_upload": bool(LOT_RULES.get("block_upload_without_lot", True)),
            "excel": {"file": EXCEL.path().name if EXCEL.folder else "", "updated_at": EXCEL.updated_at,
                      "locked": EXCEL.locked},
            "backup": HISTORY.last,
            "devices": devices, "recent": recent,
            "pending": pending_count(), "now": now_iso(),
            "upload_folder": UPLOAD_CFG.get("folder", ""),
            "history": load_upload_state().get("history", [])[:8]}


def device_state(name):
    with DEV_LOCK:
        d = dict(DEVICES.get(name) or {})
        lv = LIVE.get(name) or {}
        lines = list(lv.get("lines") or [])[-60:]
        test = json.loads(json.dumps(lv.get("test"))) if lv.get("test") else None
        results = [r for r in RECENT if r["instrument"] == name][:30]
    return {"device": d, "lines": lines, "test": test, "results": results}


DASHBOARD_HTML = r"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Lab Logger</title>
<link href="https://fonts.googleapis.com/css2?family=IBM+Plex+Mono:wght@400;500&family=IBM+Plex+Sans:wght@400;500;600&display=swap" rel="stylesheet">
<style>
:root{--bg:#F3F4F6;--card:#FFFFFF;--line:#E2E5EA;--line2:#EEF0F3;--ink:#1B2430;--muted:#5B6573;--accent:#1F5FBF;--accent2:#174A94;
--ok:#1E8E3E;--okHalo:#DDF1E2;--okInk:#17692F;--bad:#C5221F;--badHalo:#F9DEDC;--badInk:#A11B18;--warnBg:#FDF0DC;--warnInk:#8A4B08;--warn:#D9822B;
--sans:'IBM Plex Sans',system-ui,-apple-system,'Segoe UI',sans-serif;--mono:'IBM Plex Mono',ui-monospace,Consolas,monospace}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--ink);font-family:var(--sans)}
a{color:var(--accent)}a:hover{color:var(--accent2)}
header{display:flex;flex-wrap:wrap;align-items:center;justify-content:space-between;gap:16px;padding:16px 32px;background:var(--card);border-bottom:1px solid var(--line)}
.brand{display:flex;align-items:baseline;gap:14px;flex-wrap:wrap}.brand b{font-size:20px;font-weight:600}.muted{color:var(--muted)}
main{max-width:1216px;margin:0 auto;padding:28px 32px 48px;display:flex;flex-direction:column;gap:24px}
.card{background:var(--card);border:1px solid var(--line);border-radius:12px}
.btn{display:inline-flex;align-items:center;justify-content:center;gap:8px;min-height:44px;padding:0 18px;border-radius:8px;border:1px solid var(--line);background:#fff;color:var(--ink);font:500 15px var(--sans);cursor:pointer;text-decoration:none}
.btn.primary{background:var(--accent);border-color:var(--accent);color:#fff}.btn.primary:hover{background:var(--accent2);color:#fff}
.btn.dark{background:var(--ink);border-color:var(--ink);color:#fff}.btn:disabled{opacity:.5;cursor:not-allowed}
.ctx{display:flex;flex-wrap:wrap;align-items:flex-end;gap:16px;padding:18px 20px}
.field{display:flex;flex-direction:column;gap:6px}.field label{font-size:13px;font-weight:500;color:var(--muted)}
input[type=date]{min-width:165px}
input[type=text],input[type=date],select{min-height:44px;padding:0 12px;border:1px solid #C9CED6;border-radius:8px;font:15px var(--sans);color:var(--ink);background:#fff}
.grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(250px,1fr));gap:16px}
.dev{display:flex;flex-direction:column;gap:12px;padding:20px;text-decoration:none;color:var(--ink)}.dev:hover{border-color:#B8C2D0;color:var(--ink)}
.row{display:flex;align-items:center;justify-content:space-between;gap:12px}
.status{display:inline-flex;align-items:center;gap:8px;font-size:13px;font-weight:500}
.dot{width:12px;height:12px;border-radius:50%;flex:none}
.ok .dot{background:var(--ok);box-shadow:0 0 0 4px var(--okHalo)}.ok{color:var(--okInk)}
.bad .dot{background:var(--bad);box-shadow:0 0 0 4px var(--badHalo)}.bad{color:var(--badInk)}
.big{font-family:var(--mono);font-size:34px;font-weight:500;letter-spacing:-.02em}
.foot span:last-child{white-space:nowrap}
.foot{display:flex;justify-content:space-between;gap:8px;font-size:13px;color:var(--muted);border-top:1px solid var(--line2);padding-top:12px}
.pill{display:inline-flex;align-items:center;gap:8px;padding:6px 12px;border-radius:999px;background:var(--warnBg);color:var(--warnInk);font-size:13px;font-weight:500}
.pill .dot{width:8px;height:8px;background:var(--warn)}
table{width:100%;border-collapse:collapse;font-size:14px;min-width:640px}
th{text-align:left;color:var(--muted);font-size:12px;text-transform:uppercase;letter-spacing:.04em;font-weight:500;padding:10px 12px}
td{padding:11px 12px;border-top:1px solid var(--line2)}td.m,th.r,td.r{font-family:var(--mono)}td.r,th.r{text-align:right}
.scroll{overflow-x:auto}.hd{display:flex;flex-wrap:wrap;justify-content:space-between;align-items:center;gap:12px;padding:16px 20px;border-bottom:1px solid var(--line)}
h2{margin:0;font-size:16px;font-weight:600}h3{margin:0;font-size:13px;font-weight:600;color:var(--muted);text-transform:uppercase;letter-spacing:.04em}
.split{display:flex;flex-wrap:wrap;gap:20px;align-items:flex-start}.split>.mainc{flex:999 1 560px;min-width:0}.split>.side{flex:1 1 300px;display:flex;flex-direction:column;gap:16px}
.pad{padding:24px;display:flex;flex-direction:column;gap:20px}
pre{margin:0;padding:14px 16px;background:#1B2430;color:#E6E9EE;border-radius:8px;font:13px/1.6 var(--mono);overflow:auto;max-height:280px}
.kv{display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:20px;flex:1 1 360px}.kv span:first-child{display:block;font-size:12px;color:var(--muted)}.kv span:last-child{font-family:var(--mono);font-size:18px}
.list>div{display:flex;justify-content:space-between;gap:12px;padding:10px 0;border-top:1px solid var(--line2)}
fieldset{margin:0;padding:0;border:0;display:flex;flex-direction:column;gap:6px}legend{padding:0;margin-bottom:8px;font-size:16px;font-weight:600}
.opt{display:flex;align-items:center;gap:10px;min-height:44px;font-size:15px}.opt input{width:18px;height:18px}
.msg{padding:12px 16px;border-radius:8px;font-size:14px}.msg.good{background:var(--okHalo);color:var(--okInk)}.msg.err{background:var(--badHalo);color:var(--badInk)}
.back{display:inline-flex;align-items:center;gap:6px;min-height:44px;font-size:15px;font-weight:500;text-decoration:none}
@media (max-width:640px){header,main{padding-left:16px;padding-right:16px}.kv{grid-template-columns:repeat(2,minmax(0,1fr))}}
</style>
</head>
<body>
<div id="app"><main><p class="muted">Loading…</p></main></div>
<script>
const $ = (s, el=document) => el.querySelector(s);
const esc = v => String(v ?? '').replace(/[&<>"']/g, c => ({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[c]));
let S = null, D = null, uploadMsg = '';
const post = (path, body) => fetch(path, {method:'POST', headers:{'Content-Type':'application/json','X-Lab-Logger':'1'}, body: JSON.stringify(body||{})});
const hhmm = t => t ? String(t).slice(11,16) : '';
const ago = t => { if(!t) return 'never'; const s = Math.round((Date.now() - new Date(t).getTime())/1000);
  return s < 60 ? s + ' s ago' : s < 3600 ? Math.round(s/60) + ' min ago' : hhmm(t); };
const statusHtml = d => d.ok ? '<span class="status ok"><span class="dot" aria-hidden="true"></span>Connected</span>'
                            : '<span class="status bad"><span class="dot" aria-hidden="true"></span>Not connected</span>';
// What an operator should read for each instrument, in plain words.
function plain(d){
  const sup = esc(S.support || 'your supervisor');
  const det = (d.detail || '').toLowerCase();
  if(d.type === 'mb90' || d.type === 'ph'){
    return d.ok ? {short:'Plugged in and ready', help: d.type === 'mb90' ? 'Run tests as normal. Results save by themselves when each test finishes.' : 'Take a reading, then tap the printer icon on the meter to save it.'}
                : {short:'Check the USB cable', help:'Make sure the instrument is switched on and its USB cable is plugged into this laptop. It reconnects by itself within a few seconds. Still red? Tell ' + sup + '.'};
  }
  if(d.type === 'colorimeter_folder'){
    return d.ok ? {short:'Ready for exports', help:'After scanning, click Export Excel in the Linshang software and save into the Colorimeter Exports folder.'}
                : {short:'Export folder missing', help:'Tell ' + sup + '.'};
  }
  if(d.type === 'skala_api' || d.type === 'export_folder'){
    if(d.ok) return {short:'Linked to SKALA', help:'Run tests as normal in SKALA. Readings appear here a few minutes after each test.'};
    if(det.includes('token') || det.includes('sign')) return {short:'SKALA sign-in expired', help:'Readings are safe in SKALA, but they can\'t come through until the sign-in is renewed. Tell ' + sup + '.'};
    return {short:'Waiting for internet', help:'Nothing to do. Readings are safe in SKALA and come through by themselves when this laptop is back online.'};
  }
  return {short: d.detail || '', help: ''};
}
const icon = {up:'<svg width="18" height="18" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round" aria-hidden="true"><path d="M12 16V4"/><path d="M7 9l5-5 5 5"/><path d="M4 20h16"/></svg>',
  back:'<svg width="18" height="18" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round" aria-hidden="true"><path d="M15 18l-6-6 6-6"/></svg>'};

function header(inner){ return `<header>${inner}</header>`; }
function topbar(){
  return header(`<div class="brand"><b>Lab Logger</b><span class="muted">${esc(S.station)}</span></div>
   <div class="row" style="flex-wrap:wrap"><span class="muted" style="font-size:14px" id="pending">${pendingText()}</span>
   <a class="btn primary" href="#upload">${icon.up} Upload data</a></div>`);
}
const pendingText = () => `${S.pending} result${S.pending==1?'':'s'} not yet uploaded`;

// The home page is drawn once; after that only the cards, results and messages are
// updated, so the Lot / Operator / Note boxes are never redrawn while someone types.
function homeShell(){
  const c = S.context;
  return topbar() + `<main><div id="banner"></div>
    <form class="card ctx" id="ctx" aria-label="Current lot and operator" autocomplete="off">
      <div class="field"><label for="lot">Lot</label><input type="text" id="lot" value="${esc(c.lot)}" placeholder="${esc(S.lot_example || 'Lot number')}" style="width:220px;font-family:var(--mono)"></div>
      <div class="field"><label for="op">Operator</label>${operatorInput('op', c.operator)}</div>
      <div class="field" style="flex:1 1 220px"><label for="note">Note (optional)</label><input type="text" id="note" value="${esc(c.note)}" placeholder="e.g. re-test" style="width:100%"></div>
      <div style="display:flex;flex-direction:column;gap:6px;align-items:flex-start"><span id="saveState" style="font-size:13px;min-height:18px"></span>
        <button class="btn dark" type="submit">Save</button></div>
      <div id="ctxInfo" style="flex:1 1 100%;display:flex;flex-wrap:wrap;align-items:center;gap:12px;font-size:14px"></div>
    </form>
    <section class="card" id="needs" aria-label="Results that need a lot" style="display:none;border-color:#E3A29D"></section>
    <section class="grid" aria-label="Instruments" id="cards"></section>
    <section class="card" aria-label="Recent results"><div class="hd"><h2>Recent results</h2>
        <div class="row" style="flex-wrap:wrap;justify-content:flex-end"><span id="excelInfo" class="muted" style="font-size:13px"></span>
        <button class="btn" type="button" id="openExcel">Open in Excel</button></div></div>
      <div class="scroll"><table><thead><tr><th>Time</th><th>Lot</th><th>Instrument</th><th>Measurement</th><th class="r">Value</th><th>Operator</th></tr></thead><tbody id="recentBody"></tbody></table></div></section>
  </main>`;
}
const FIELDS = {lot: 'lot', op: 'operator', note: 'note'};
function operatorInput(id, value){
  if(S.operators && S.operators.length){
    return `<select id="${id}" style="width:220px"><option value="">Choose your name</option>${S.operators.map(o => `<option ${o === value ? 'selected' : ''}>${esc(o)}</option>`).join('')}</select>`;
  }
  return `<input type="text" id="${id}" value="${esc(value)}" placeholder="Your name" style="width:220px">`;
}
let needsKey = null, needsMsg = '', excelMsgAt = 0;
let justSaved = 0, saveError = 0;
function renderNeeds(){
  const box = $('#needs');
  const list = S.needs_lot || [];
  const key = list.map(r => r.record_id).join('|') + '#' + needsMsg;
  if(key === needsKey) return;          // only redraw when the list changes, so typing isn't lost
  const keepLot = $('#nLot') ? $('#nLot').value : (S.context.lot || '');
  const keepOp = $('#nOp') ? $('#nOp').value : (S.context.operator || '');
  const unticked = new Set([...document.querySelectorAll('#needs input[name=nid]:not(:checked)')].map(x => x.value));
  needsKey = key;
  if(!list.length && !needsMsg){ box.style.display = 'none'; box.innerHTML = ''; return; }
  box.style.display = '';
  const labelOf = n => (S.devices.find(x => x.name === n) || {}).label || n;
  const rows = list.map(r => `<label class="opt" style="border-top:1px solid var(--line2);padding:4px 0;gap:14px">
      <input type="checkbox" name="nid" value="${esc(r.record_id)}" ${unticked.has(r.record_id) ? '' : 'checked'}>
      <span class="m muted" style="width:52px;font-family:var(--mono)">${esc(hhmm(r.time))}</span><span style="flex:1 1 160px">${esc(labelOf(r.instrument))}</span>
      <span style="flex:1 1 140px">${esc(r.measurement)}</span><span style="font-family:var(--mono);font-weight:500">${esc(r.value)} <span class="muted">${esc(r.unit)}</span></span></label>`).join('');
  box.innerHTML = `<div class="pad" style="gap:14px">
     <div style="display:flex;flex-direction:column;gap:4px"><h2 style="color:var(--badInk)">${list.length ? list.length + ' result' + (list.length==1?'':'s') + ' need a lot number' : 'All results have a lot number'}</h2>
     ${list.length ? '<span class="muted" style="font-size:14px">These finished while no lot was entered. Tick the ones from the same lot, enter the lot and your name, then press Assign.</span>' : ''}</div>
     ${needsMsg}
     ${list.length ? `<div>${rows}</div>
     <form id="assign" style="display:flex;flex-wrap:wrap;align-items:flex-end;gap:12px">
       <div class="field"><label for="nLot">Lot</label><input type="text" id="nLot" value="${esc(keepLot)}" placeholder="${esc(S.lot_example || 'Lot number')}" style="width:200px;font-family:var(--mono)"></div>
       <div class="field"><label for="nOp">Your name</label>${operatorInput('nOp', keepOp)}</div>
       <button class="btn primary" type="submit">Assign to ticked results</button></form>` : ''}</div>`;
  const f = $('#assign');
  if(f) f.onsubmit = async e => { e.preventDefault();
    const ids = [...document.querySelectorAll('#needs input[name=nid]:checked')].map(x => x.value);
    if(!ids.length){ needsMsg = '<div class="msg err">Tick at least one result.</div>'; needsKey = null; renderNeeds(); return; }
    const res = await post('/api/assign', {record_ids: ids, lot: $('#nLot').value.trim(), operator: $('#nOp').value.trim()});
    const j = await res.json();
    needsMsg = j.ok ? `<div class="msg good">✓ ${j.assigned} result${j.assigned==1?'':'s'} assigned.</div>` : `<div class="msg err">${esc(j.error)}</div>`;
    needsKey = null; await refresh();
    if(j.ok) setTimeout(() => { needsMsg = ''; needsKey = null; renderNeeds(); }, 4000);
  };
}
function updateHome(){
  const c = S.context;
  $('#pending').textContent = pendingText();
  // Keep the boxes in step with what's saved, except a box someone has changed and not saved yet.
  let unsaved = false;
  for(const [id, key] of Object.entries(FIELDS)){
    const el = $('#' + id);
    if(el.dataset.dirty === '1'){ if(el.value === (c[key] || '')) el.dataset.dirty = ''; else unsaved = true; }
    else if(document.activeElement !== el && el.value !== (c[key] || '')) el.value = c[key] || '';
  }
  const st = $('#saveState');
  if(Date.now() - saveError < 6000){ /* keep showing the problem with what was typed */ }
  else if(unsaved){ st.textContent = 'Not saved yet: press Save'; st.style.color = 'var(--warnInk)'; }
  else if(Date.now() - justSaved < 4000){ st.textContent = '✓ Saved'; st.style.color = 'var(--okInk)'; }
  else st.textContent = '';
  const info = $('#ctxInfo');
  if(c.lot && c.operator){
    const until = S.expires_at ? ' It clears by itself at ' + esc(hhmm(S.expires_at)) + (S.expires_at.slice(0,10) !== S.now.slice(0,10) ? ' (midnight)' : '') + '.' : '';
    const html = `<span class="muted">Results are being labelled <b style="color:var(--ink);font-family:var(--mono);font-weight:500">${esc(c.lot)}</b> · <b style="color:var(--ink);font-weight:500">${esc(c.operator)}</b>.${until}</span>
      <button class="btn" type="button" id="finishLot">Finished this lot</button>`;
    if(info.dataset.html !== html){ info.innerHTML = html; info.dataset.html = html;
      $('#finishLot').onclick = async () => { await post('/api/context', {finish: true}); ['lot','op','note'].forEach(id => { $('#'+id).value=''; $('#'+id).dataset.dirty=''; }); await refresh(); }; }
  } else if(info.dataset.html !== ''){ info.innerHTML = ''; info.dataset.html = ''; }
  renderNeeds();
  const xi = $('#excelInfo'), x = S.excel || {};
  if(Date.now() - excelMsgAt > 5000){
    xi.style.color = x.locked ? 'var(--badInk)' : '';
    xi.textContent = x.locked ? 'The Excel file is open, so it can\'t update. Close it and use Open in Excel instead.'
                   : x.updated_at ? 'Everything is saved to Excel (updated ' + hhmm(x.updated_at) + ')' : 'Everything is saved to Excel';
  }
  $('#banner').innerHTML = (!c.lot || !c.operator)
    ? `<div class="msg err" role="status" style="font-size:15px"><b>Before testing:</b> enter the lot number and your name, then press Save. Results that finish before then will need a lot assigned.</div>` : '';
  $('#cards').innerHTML = S.devices.map(d => {
    const r = d.last_result;
    return `<a class="card dev" href="#device/${encodeURIComponent(d.name)}">
      <div class="row"><span style="font-size:16px;font-weight:600">${esc(d.label)}</span>${statusHtml(d)}</div>
      <div class="muted" style="font-size:13px" title="${esc(d.detail)}">${esc(plain(d).short)}${d.live_active ? ' · <b style="color:var(--warnInk)">test running</b>' : ''}</div>
      <div style="display:flex;align-items:baseline;gap:8px"><span class="big">${r ? esc(r.value) : '—'}</span><span class="muted">${r ? esc(r.unit) : ''}</span></div>
      <div class="foot"><span>${r ? 'Last result ' + esc(hhmm(r.time)) + (r.lot ? ' · ' + esc(r.lot) : '') : 'No results yet'}</span><span style="color:var(--accent);font-weight:500">Open →</span></div></a>`;
  }).join('');
  const labelOf = n => (S.devices.find(x => x.name === n) || {}).label || n;
  $('#recentBody').innerHTML = S.recent.map(r => `<tr><td class="m muted">${esc(hhmm(r.time))}</td><td class="m">${esc(r.lot||'—')}</td><td>${esc(labelOf(r.instrument))}</td>
     <td>${esc(r.measurement)}</td><td class="r"><b style="font-weight:500">${esc(r.value)}</b> <span class="muted">${esc(r.unit)}</span></td><td>${esc(r.operator)}</td></tr>`).join('')
     || '<tr><td colspan="6" class="muted" style="padding:20px">No results yet today. They appear here by themselves as soon as an instrument finishes a test.</td></tr>';
}

function secs(mmss){ const [m,s] = String(mmss).split(':').map(Number); return (m||0)*60 + (s||0); }
function curveSvg(points){
  if(!points || points.length < 2) return '<p class="muted" style="margin:0">The drying curve appears here once the MB90 sends interval readings (turn on Print settings → Intermediate results print interval).</p>';
  const xs = points.map(p => secs(p[0])), ys = points.map(p => parseFloat(p[1]) || 0);
  const xmax = Math.max(120, Math.max(...xs) * 1.15), ymax = Math.max(1, Math.max(...ys) * 1.15);
  const X = x => 48 + x / xmax * 582, Y = y => 210 - y / ymax * 190;
  const pts = xs.map((x,i) => `${X(x).toFixed(1)},${Y(ys[i]).toFixed(1)}`).join(' ');
  let grid = '', labels = '';
  for(let i=0;i<=4;i++){ const v = ymax/4*i; grid += `<line x1="48" y1="${Y(v)}" x2="630" y2="${Y(v)}" stroke="#EEF0F3"/>`;
    labels += `<text x="40" y="${Y(v)+4}" text-anchor="end">${v.toFixed(1)}</text>`; }
  for(let i=0;i<=3;i++){ const t = xmax/3*i; labels += `<text x="${X(t)}" y="230" text-anchor="middle">${Math.floor(t/60)}:${String(Math.round(t%60)).padStart(2,'0')}</text>`; }
  const last = points[points.length-1];
  return `<svg viewBox="0 0 640 240" width="100%" height="240" role="img" aria-label="Drying curve, latest ${esc(last[1])} at ${esc(last[0])}">
    ${grid}<line x1="48" y1="210" x2="630" y2="210" stroke="#C9CED6"/><g font-family="IBM Plex Mono, monospace" font-size="11" fill="#5B6573">${labels}</g>
    <polyline fill="none" stroke="#1F5FBF" stroke-width="2.5" stroke-linejoin="round" stroke-linecap="round" points="${pts}"/>
    <circle cx="${X(xs[xs.length-1])}" cy="${Y(ys[ys.length-1])}" r="5" fill="#1F5FBF"/></svg>`;
}

function viewDevice(name){
  if(!D || !D.device || D.device.name !== name) return topbar() + '<main><p class="muted">Loading…</p></main>';
  const d = D.device, t = D.test, last = D.results[0];
  let livePanel = '';
  if(d.type === 'mb90' && t){
    livePanel = `<div class="row" style="flex-wrap:wrap"><h2>${t.active ? 'Live test' : 'Last test'}</h2>
        ${t.active ? '<span class="pill"><span class="dot" aria-hidden="true"></span>Drying · live</span>' : '<span class="muted" style="font-size:13px">Finished ' + esc(hhmm(t.finished)) + '</span>'}</div>
      <div style="display:flex;flex-wrap:wrap;align-items:flex-end;gap:32px">
        <div style="display:flex;align-items:baseline;gap:10px"><span class="big" style="font-size:64px;line-height:1">${esc(t.value||'—')}</span><span class="muted" style="font-size:20px">${esc(t.unit||'')}</span></div>
        <div class="kv"><div><span>Elapsed</span><span>${esc(t.elapsed||'—')}</span></div><div><span>Drying temp</span><span>${t.temp ? esc(t.temp)+' °C' : '—'}</span></div>
        <div><span>Start weight</span><span>${t.start_weight ? esc(t.start_weight)+' g' : '—'}</span></div><div><span>Method</span><span style="font-family:var(--sans);font-size:16px">${esc(t.method||'—')}</span></div></div></div>
      <figure style="margin:0">${curveSvg(t.points)}</figure>`;
  } else {
    livePanel = `<div class="row"><h2>Latest result</h2>${last ? '<span class="muted" style="font-size:13px">' + esc(last.time.replace('T',' ').slice(0,16)) + '</span>' : ''}</div>
      <div style="display:flex;align-items:baseline;gap:10px"><span class="big" style="font-size:64px;line-height:1">${last ? esc(last.value) : '—'}</span><span class="muted" style="font-size:20px">${last ? esc(last.unit) : ''}</span></div>
      <p class="muted" style="margin:0">${last ? esc(last.measurement) + (last.lot ? ' · lot ' + esc(last.lot) : '') + (last.operator ? ' · ' + esc(last.operator) : '') : hint(d)}</p>`;
  }
  const lines = D.lines.length ? esc(D.lines.map(l => l[1]).join('\n')) : '';
  const feed = (d.type === 'mb90' || d.type === 'ph')
    ? `<details ${S._openFeed ? 'open' : ''} id="feedBox"><summary style="cursor:pointer;min-height:44px;display:flex;align-items:center;font-size:14px;color:var(--muted)">Show what the instrument printed</summary>${lines ? '<pre id="feed">' + lines + '</pre>' : '<p class="muted" style="margin:0">Nothing received yet.</p>'}</details>` : '';
  const results = D.results.map(r => `<div><div style="display:flex;flex-direction:column;gap:2px"><span style="font-size:14px;font-weight:500">${esc(r.measurement)}</span>
      <span class="muted" style="font-size:13px">${esc(hhmm(r.time))} · ${esc(r.lot||'no lot')}</span></div><span style="font-family:var(--mono);font-weight:500">${esc(r.value)} ${esc(r.unit)}</span></div>`).join('')
      || '<div><span class="muted">No results yet</span></div>';
  const skala = d.type === 'skala_api' ? '<button class="btn" id="sync" type="button">Check SKALA now</button>' : '';
  return header(`<div class="row" style="flex-wrap:wrap;justify-content:flex-start;gap:16px"><a class="back" href="#">${icon.back} All instruments</a>
      <b style="font-size:20px">${esc(d.label)}</b>${statusHtml(d)}<span class="muted" style="font-size:14px">${esc(plain(d).short)}</span></div>
      <span class="muted" style="font-size:14px">Lot <b style="font-family:var(--mono);color:var(--ink);font-weight:500">${esc(S.context.lot||'—')}</b> · ${esc(S.context.operator||'no operator')}</span>`)
   + `<main><div class="split"><section class="card pad mainc" aria-label="Live">${livePanel}${feed}</section>
      <aside class="side"><section class="card pad" style="gap:4px"><h2 style="margin-bottom:8px">Recent</h2><div class="list">${results}</div></section>
      <section class="card pad" style="gap:12px"><h2>${d.ok ? 'All good' : 'What to do'}</h2><p style="margin:0;font-size:15px;line-height:1.5">${plain(d).help}</p>${skala}
        <details id="supBox"><summary style="cursor:pointer;min-height:44px;display:flex;align-items:center;font-size:13px;color:var(--muted)">For support</summary>
        <div class="muted" style="font-size:13px;line-height:1.6">${esc(d.detail)}<br>Last data received: ${esc(ago(d.last_data))}</div></details></section></aside></div></main>`;
}
function hint(d){
  return {mb90:'Run a test; results appear when it finishes.', ph:'Take a reading and tap the printer icon on the meter.',
    colorimeter_folder:'Export Excel from the Linshang software into the Colorimeter Exports folder.',
    skala_api:'Readings come through from SKALA a few minutes after each test, whenever the laptop is online.', export_folder:'Save SKALA downloads into the watch folder.'}[d.type] || '';
}

function viewUpload(){
  const lots = [...new Set(S.recent.map(r => r.lot).filter(Boolean))];
  const folderOk = !!S.upload_folder;
  const blocked = S.block_upload && S.needs_lot_count > 0;
  const hist = S.history.map(h => `<div style="flex-direction:column;align-items:flex-start;gap:2px"><span style="font-size:14px;font-weight:500">${esc(h.when.replace('T',' ').slice(0,16))}</span>
     <span class="muted" style="font-size:13px">${esc(h.results)} results · ${esc(h.scope)} · ${esc(h.dest)}${h.operator ? ' · ' + esc(h.operator) : ''}</span></div>`).join('') || '<div><span class="muted">No uploads yet</span></div>';
  return header(`<div class="row" style="justify-content:flex-start;gap:16px;flex-wrap:wrap"><a class="back" href="#">${icon.back} All instruments</a><b style="font-size:20px">Upload data</b></div>`)
  + `<main><div class="split"><form class="card pad mainc" id="up">
     ${uploadMsg}
     <div style="display:flex;flex-direction:column;gap:6px"><h2 style="font-size:22px">${S.pending ? S.pending + ' result' + (S.pending==1?'':'s') + ' not uploaded yet' : 'Everything is uploaded'}</h2>
       <p class="muted" style="margin:0;font-size:15px">${folderOk ? 'Upload now copies them to the shared cloud folder.' : 'Upload now saves an Excel file of them to this laptop.'}</p></div>
     ${blocked ? `<div class="msg err" style="font-size:15px">${S.needs_lot_count} result${S.needs_lot_count==1?'':'s'} still need a lot number. <a href="#">Assign them on the main page</a>, then come back to upload.</div>` : ''}
     <div><button class="btn primary" type="submit" style="min-height:56px;padding:0 32px;font-size:18px" ${S.pending && !blocked ? '' : 'disabled'}>${icon.up} Upload now</button></div>
     <details id="more"><summary style="cursor:pointer;min-height:44px;display:flex;align-items:center;font-size:14px;color:var(--muted)">More options (one lot, a date range, other files)</summary>
     <div style="display:flex;flex-direction:column;gap:16px;padding-top:8px">
     <fieldset><legend>What to upload</legend>
       <label class="opt"><input type="radio" name="scope" value="new" checked> Everything not uploaded yet <span class="muted">(${S.pending} results)</span></label>
       <label class="opt"><input type="radio" name="scope" value="lot"> One lot
         <select id="lotSel" aria-label="Lot">${lots.map(l => `<option>${esc(l)}</option>`).join('') || '<option value="">no lots yet</option>'}</select></label>
       <label class="opt" style="flex-wrap:wrap"><input type="radio" name="scope" value="range"> A date range
         <input type="date" id="from" aria-label="From date"> to <input type="date" id="to" aria-label="To date"></label></fieldset>
     <fieldset style="border-top:1px solid var(--line2);padding-top:16px"><legend>Files</legend>
       <label class="opt"><input type="checkbox" name="inc" value="shared" checked> Shared: All Results and Lot Summary (Excel)</label>
       <label class="opt"><input type="checkbox" name="inc" value="detail" checked> Each instrument's full detail files</label>
       <label class="opt"><input type="checkbox" name="inc" value="raw"> Raw instrument text (audit copies)</label></fieldset>
     <fieldset style="border-top:1px solid var(--line2);padding-top:16px"><legend>Where to</legend>
       <label class="opt"><input type="radio" name="dest" value="folder" ${folderOk ? 'checked' : 'disabled'}> Shared cloud folder
         <span class="muted" style="font-size:13px">${folderOk ? '' : '(not set up yet, ask ' + esc(S.support) + ')'}</span></label>
       <label class="opt"><input type="radio" name="dest" value="download" ${folderOk ? '' : 'checked'}> Download to this laptop</label></fieldset>
     <div><button class="btn" type="submit" id="upMore" style="min-height:48px">${icon.up} Upload with these options</button></div>
     </div></details>
     <p class="muted" style="margin:0;font-size:14px">No internet? That's fine. The cloud folder sends everything once the laptop is back online.</p>
   </form>
   <aside class="side"><section class="card pad" style="gap:4px"><h2 style="margin-bottom:8px">Recent uploads</h2><div class="list">${hist}</div></section></aside></div></main>`;
}

function render(){
  if(!S) return;
  const h = location.hash.replace(/^#\/?/, '');
  const app = $('#app');
  if(h === ''){                           // home: draw once, then only update around the form
    if(!$('#ctx')){ app.innerHTML = homeShell(); bind(); }
    updateHome();
    return;
  }
  const openIds = [...document.querySelectorAll('details[open]')].map(x => x.id).filter(Boolean);
  if(h.startsWith('device/')) app.innerHTML = viewDevice(decodeURIComponent(h.slice(7)));
  else if(h === 'upload') { if(!$('#up')) app.innerHTML = viewUpload(); }
  openIds.forEach(id => { const x = document.getElementById(id); if(x) x.open = true; });
  const feed = $('#feed'); if(feed) feed.scrollTop = feed.scrollHeight;
  const fb = $('#feedBox'); if(fb) fb.ontoggle = () => { S._openFeed = fb.open; };
  bind();
}
function bind(){
  const ctx = $('#ctx');
  if(ctx){
    ctx.querySelectorAll('input,select').forEach(i => { i.oninput = i.onchange = () => { i.dataset.dirty = '1'; saveError = 0; updateHome(); }; });
    ctx.onsubmit = async e => { e.preventDefault();
      const body = {lot: $('#lot').value.trim(), operator: $('#op').value.trim(), note: $('#note').value.trim()};
      const res = await post('/api/context', body); const j = await res.json();
      if(!j.ok){ const st = $('#saveState'); st.textContent = j.error; st.style.color = 'var(--badInk)'; saveError = Date.now(); return; }
      ['lot','op','note'].forEach(id => { $('#' + id).dataset.dirty = ''; });
      justSaved = Date.now(); await refresh(); };
  }
  const ox = $('#openExcel');
  if(ox) ox.onclick = async () => { ox.disabled = true; const xi = $('#excelInfo'); xi.textContent = 'Opening Excel…'; xi.style.color = '';
    excelMsgAt = Date.now();
    try { const j = await (await post('/api/open_excel')).json();
      xi.textContent = j.ok ? 'Opened in Excel. This is a copy, so the logger keeps saving while it\'s open.' : j.error; xi.style.color = j.ok ? 'var(--okInk)' : 'var(--badInk)'; }
    catch(e){ xi.textContent = 'Excel didn\'t open. Try again.'; }
    excelMsgAt = Date.now(); ox.disabled = false; };
  const sync = $('#sync'); if(sync) sync.onclick = async () => { sync.disabled = true; sync.textContent = 'Checking…'; await post('/api/sync'); };
  const up = $('#up');
  if(up) up.onsubmit = async e => {
    e.preventDefault();
    const f = new FormData(up);
    const useOptions = e.submitter && e.submitter.id === 'upMore';
    const body = useOptions
      ? {scope: f.get('scope'), lot: $('#lotSel').value, from: $('#from').value, to: $('#to').value, include: f.getAll('inc'), dest: f.get('dest')}
      : {scope: 'new', include: ['shared', 'detail'], dest: S.upload_folder ? 'folder' : 'download'};
    const btn = e.submitter || up.querySelector('button[type=submit]'); btn.disabled = true; btn.textContent = 'Uploading…';
    try {
      const res = await post('/api/upload', body);
      const type = res.headers.get('Content-Type') || '';
      if(type.includes('json')){
        const j = await res.json();
        uploadMsg = j.ok ? `<div class="msg good" style="font-size:15px">Done. ${j.results} result${j.results==1?'':'s'} uploaded.</div>` : `<div class="msg err" style="font-size:15px">${esc(j.error)}</div>`;
      } else {
        const blob = await res.blob(); const name = (res.headers.get('Content-Disposition')||'').match(/filename="(.+)"/);
        const a = document.createElement('a'); a.href = URL.createObjectURL(blob); a.download = name ? name[1] : 'lab_data';
        document.body.appendChild(a); a.click(); a.remove();
        const info = JSON.parse(res.headers.get('X-Upload-Info') || '{}');
        uploadMsg = `<div class="msg good" style="font-size:15px">Done. ${info.results ?? ''} results saved to this laptop's Downloads folder.</div>`;
      }
    } catch(err){ uploadMsg = `<div class="msg err" style="font-size:15px">The upload didn't work. Try again in a minute; if it keeps failing, tell ${esc(S.support)}.</div>`; }
    await refresh(); $('#app').innerHTML = viewUpload(); bind();
  };
}
async function refresh(){
  try {
    const keep = S && S._openFeed; S = await (await fetch('/api/state')).json(); S._openFeed = keep;
    const h = location.hash.replace(/^#\/?/, '');
    if(h.startsWith('device/')) D = await (await fetch('/api/device/' + h.slice(7))).json();
    render();
  } catch(e){ $('#app').innerHTML = '<main><div class="card pad"><h2>Lab Logger isn\'t running</h2><p style="margin:0;font-size:15px">Double-click the <b>Lab Logger</b> icon on the desktop. This page comes back by itself.</p></div></main>'; }
}
window.addEventListener('hashchange', () => { uploadMsg = ''; D = null; needsKey = null; needsMsg = ''; $('#app').innerHTML = ''; refresh(); });
refresh(); setInterval(refresh, 2000);
</script>
</body>
</html>
"""


def start_dashboard(cfg, out_dir):
    from http.server import ThreadingHTTPServer, BaseHTTPRequestHandler
    dc = cfg.get("dashboard", {})
    if not dc.get("enabled", True):
        return
    port = int(dc.get("port", 8080))
    DASH.update(out_dir=out_dir, station=cfg.get("station_name", ""), port=port,
                support=dc.get("support_contact") or "your supervisor")
    UPLOAD_CFG.update(cfg.get("upload", {}))
    allowed_hosts = {f"localhost:{port}", f"127.0.0.1:{port}", "localhost", "127.0.0.1"}

    class Handler(BaseHTTPRequestHandler):
        def log_message(self, *a):
            pass

        def _host_ok(self):
            return (self.headers.get("Host") or "") in allowed_hosts

        def _send(self, code, body, ctype="application/json", extra=None):
            data = body if isinstance(body, bytes) else body.encode("utf-8")
            self.send_response(code)
            self.send_header("Content-Type", ctype)
            self.send_header("Content-Length", str(len(data)))
            self.send_header("Cache-Control", "no-store")
            for k, v in (extra or {}).items():
                self.send_header(k, v)
            self.end_headers()
            self.wfile.write(data)

        def _json(self, obj, code=200):
            self._send(code, json.dumps(obj, default=str))

        def do_GET(self):
            if not self._host_ok():
                return self._send(403, "forbidden", "text/plain")
            path = self.path.split("?")[0]
            if path in ("/", "/index.html"):
                return self._send(200, DASHBOARD_HTML, "text/html; charset=utf-8")
            if path == "/api/state":
                return self._json(dashboard_state())
            if path.startswith("/api/device/"):
                import urllib.parse
                return self._json(device_state(urllib.parse.unquote(path[len("/api/device/"):])))
            self._send(404, "not found", "text/plain")

        def do_POST(self):
            if not self._host_ok() or self.headers.get("X-Lab-Logger") != "1":
                return self._send(403, "forbidden", "text/plain")
            n = int(self.headers.get("Content-Length") or 0)
            try:
                body = json.loads(self.rfile.read(n) or b"{}")
            except json.JSONDecodeError:
                return self._json({"ok": False, "error": "bad request"}, 400)
            path = self.path.split("?")[0]
            if path == "/api/context":
                if body.get("finish"):
                    clear_context("'Finished this lot' pressed")
                    return self._json({"ok": True, "context": dict(CONTEXT)})
                lot = str(body.get("lot", "")).strip()
                operator = str(body.get("operator", "")).strip()
                problem = check_context(lot, operator)
                if problem:
                    return self._json({"ok": False, "error": problem}, 400)
                with CONTEXT_LOCK:
                    CONTEXT.update(lot=lot, operator=operator, note=str(body.get("note", "")).strip(),
                                   set_at=now_iso())
                log(f"Dashboard: lot='{lot}' operator='{operator}' note='{CONTEXT['note']}'")
                return self._json({"ok": True, "context": dict(CONTEXT)})
            if path == "/api/assign":
                try:
                    n = assign_lot(body.get("record_ids") or [], str(body.get("lot", "")).strip(),
                                   str(body.get("operator", "")).strip())
                except ValueError as e:
                    return self._json({"ok": False, "error": str(e)}, 400)
                return self._json({"ok": True, "assigned": n})
            if path == "/api/open_excel":
                try:
                    return self._json({"ok": True, "path": EXCEL.open_copy()})
                except Exception as e:
                    return self._json({"ok": False, "error": str(e)}, 400)
            if path == "/api/shutdown":
                log("Stopped from Stop Lab Logger.")
                threading.Timer(0.5, STOP.set).start()  # reply first, then stop
                return self._json({"ok": True})
            if path == "/api/sync":
                SKALA_SYNC_NOW.set()
                return self._json({"ok": True})
            if path == "/api/upload":
                try:
                    info, sent = do_upload(body)
                except Exception as e:
                    return self._json({"ok": False, "error": str(e)}, 400)
                if sent is None:
                    return self._json(info)
                data = Path(sent).read_bytes()
                ctype = ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                         if str(sent).endswith(".xlsx") else "application/zip")
                return self._send(200, data, ctype, {
                    "Content-Disposition": f'attachment; filename="{Path(sent).name}"',
                    "X-Upload-Info": json.dumps(info)})
            self._send(404, "not found", "text/plain")

    try:
        srv = ThreadingHTTPServer(("127.0.0.1", port), Handler)
    except OSError as e:
        log(f"Dashboard could not start on port {port}: {e}. Is the logger already running?")
        return
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    url = f"http://localhost:{port}"
    log(f"Dashboard: {url}")
    if dc.get("open_browser", True):
        try:
            import webbrowser
            webbrowser.open(url)
        except Exception:
            pass


# --------------------------------------------------------------------------
# Config, console, entry points
# --------------------------------------------------------------------------
class StartupProblem(Exception):
    """A problem an admin can fix; shown in plain words, no traceback."""


def load_config():
    if not CONFIG_PATH.exists():
        CONFIG_PATH.write_text(json.dumps(DEFAULT_CONFIG, indent=2), encoding="utf-8")
        log(f"No settings file found, so one was made with the standard settings: {CONFIG_PATH}")
    try:
        return json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
    except json.JSONDecodeError as e:
        lines = CONFIG_PATH.read_text(encoding="utf-8", errors="replace").splitlines()
        bad = lines[e.lineno - 1].strip() if 0 < e.lineno <= len(lines) else ""
        raise StartupProblem(f"The settings file has a typo on line {e.lineno}, or the line just above it ({e.msg}): {bad}  "
                             "Usually a missing comma or quote, or a single backslash in a path "
                             "(use forward slashes, e.g. C:/LabLogger/data).")


def console_loop():
    print("Commands: lot <no.> | op <name> | note <text> | sync | status | quit")
    while not STOP.is_set():
        try:
            line = input().strip()
        except EOFError:
            return  # no keyboard (started hidden): keep logging, dashboard still works
        except KeyboardInterrupt:
            STOP.set()
            break
        cmd, _, arg = line.partition(" ")
        cmd = cmd.lower()
        if cmd in ("lot", "batch"):
            with CONTEXT_LOCK:
                CONTEXT["lot"] = arg.strip()
                CONTEXT["set_at"] = now_iso()
            log(f"Lot set to '{arg.strip()}'")
        elif cmd in ("op", "operator"):
            with CONTEXT_LOCK:
                CONTEXT["operator"] = arg.strip()
                CONTEXT["set_at"] = now_iso()
            log(f"Operator set to '{arg.strip()}'")
        elif cmd == "note":
            with CONTEXT_LOCK:
                CONTEXT["note"] = arg.strip()
            log(f"Note set to '{arg.strip()}'")
        elif cmd == "sync":
            SKALA_SYNC_NOW.set()
            log("Checking SKALA now...")
        elif cmd == "status":
            for k, v in STATUS.items():
                print(f"  {k:12s} {v}")
            print(f"  lot='{CONTEXT['lot']}' operator='{CONTEXT['operator']}' note='{CONTEXT['note']}'")
        elif cmd in ("quit", "exit", "q"):
            STOP.set()
        elif cmd:
            print("Unknown command. Use: lot, op, note, sync, status, quit")


def already_running(port):
    import socket
    try:
        with socket.create_connection(("127.0.0.1", port), timeout=1):
            return True
    except OSError:
        return False


def run_logger():
    if serial is None:
        raise StartupProblem("A required part (pyserial) isn't installed. Double-click "
                             "\"Install Lab Logger parts.bat\" in this folder, then start Lab Logger again.")
    cfg = load_config()
    port = int(cfg.get("dashboard", {}).get("port", 8080))
    if already_running(port):
        # Someone double-clicked the launcher again: just bring up the dashboard.
        import webbrowser
        webbrowser.open(f"http://localhost:{port}")
        return
    out_dir = expand(cfg.get("output_folder") or DEFAULT_CONFIG["output_folder"])
    out_dir.mkdir(parents=True, exist_ok=True)
    log(f"Lab logger started at station '{cfg.get('station_name', '')}'")
    log(f"Saving data to {out_dir}")
    SHARED.configure(cfg, out_dir)
    configure_lot_rules(cfg)
    threading.Thread(target=context_watch, daemon=True).start()
    EXCEL.configure(cfg, out_dir)
    HISTORY.configure(cfg, out_dir)
    preload_recent()

    for inst in cfg.get("instruments") or DEFAULT_CONFIG["instruments"]:
        if not inst.get("enabled", True):
            continue
        if inst["type"] in ("mb90", "ph"):
            SerialInstrument(inst, out_dir).start()
        elif inst["type"] in ("colorimeter_folder", "export_folder"):
            FolderWatcher(inst, out_dir).start()
        elif inst["type"] == "skala_api":
            SkalaFetcher(inst, out_dir).start()
        else:
            log(f"Unknown instrument type '{inst['type']}' for {inst['name']}")

    start_dashboard(cfg, out_dir)
    if sys.stdin is not None and sys.stdout is not None:
        threading.Thread(target=console_loop, daemon=True).start()
    try:
        while not STOP.is_set():
            WRITER.flush()
            housekeeping()
            STOP.wait(5)
    except KeyboardInterrupt:
        STOP.set()
    WRITER.flush()
    SHARED.rebuild_summary()
    EXCEL.rebuild_if_needed(min_gap=0)
    HISTORY.check(force=True)
    log("Logger stopped.")


def list_ports():
    if serial is None:
        sys.exit("pyserial is not installed. Run:  pip install pyserial")
    ports = list(serial.tools.list_ports.comports())
    if not ports:
        print("No serial ports found. Is the instrument plugged in and switched on?")
    for p in ports:
        print(f"{p.device:25s} {p.description}  [{p.hwid}]")


def monitor(port, baud):
    if serial is None:
        sys.exit("pyserial is not installed. Run:  pip install pyserial")
    print(f"Listening on {port} at {baud} baud. Press Print on the instrument. Ctrl+C to stop.")
    with serial.Serial(port, baud, timeout=0.5) as s:
        try:
            while True:
                data = s.read(512)
                if data:
                    sys.stdout.write(SerialInstrument.decode(data))
                    sys.stdout.flush()
        except KeyboardInterrupt:
            print("\nStopped.")


SAMPLE_MB90 = """MOISTURE DETERMINATION
Halogen Moisture Analyzer
Type MB90
SNR(Drying Unit) 1234567
SW(Drying Unit) 1.20
Method Name Method 1
Drying Prog Standard
Drying Temp 105°C
Switch Off A60(1mg/60s)
Start Weight 3.098 g
00:00 min 0.00%MC
00:30 min 9.17 %MC
01:00 min 12.35 %MC
05:21 min 31.94 %MC
Total Time. 05:21 min
End Result . 31.94 %MC
Sample ID:
1.Jan.15 15:35
---------------END----------------"""

SAMPLE_PH = """AB33PH 1234567890 1.00
0 2026/9/29 14:34 001 pH
pH 7.01 pH -2.3 mV
25.0 C ATC
98.5 % 3.2 mV"""


def selftest():
    import tempfile
    out = Path(tempfile.mkdtemp(prefix="lablogger_test_"))
    CONTEXT.update(lot="TEST-LOT", operator="selftest")
    save_serial_record("MB90", "mb90", SAMPLE_MB90, out)
    save_serial_record("pH_meter", "ph", SAMPLE_PH, out)
    colo = out / "colo_in"
    colo.mkdir()
    (colo / "export.csv").write_text("NO,Test Name,L*,a*,b*,SCAA\n1,Bean A,54.46,-25.92,15.35,96\n",
                                     encoding="utf-8")
    time.sleep(3.1)
    FolderWatcher({"name": "LS177C", "watch_folder": str(colo)}, out).import_file(colo / "export.csv")
    aq = out / "aq_in"
    aq.mkdir()
    aw = FolderWatcher({"name": "AQUALAB3", "watch_folder": str(aq), "skip_duplicates": True}, out)
    (aq / "week1.csv").write_text("Date,Product,Batch,Water Activity,Temperature\n"
                                  "2026-09-28 09:10,Roast A,B101,0.412,25.0\n"
                                  "2026-09-29 10:05,Roast A,B102,0.398,25.0\n", encoding="utf-8")
    aw.import_file(aq / "week1.csv")
    (aq / "week1_again.csv").write_text("Date,Product,Batch,Water Activity,Temperature\n"
                                        "2026-09-29 10:05,Roast A,B102,0.398,25.0\n"
                                        "2026-09-30 08:40,Roast B,B103,0.455,25.0\n", encoding="utf-8")
    aw.import_file(aq / "week1_again.csv")
    WRITER.flush()
    for f in sorted(out.rglob("*.csv")):
        print(f"\n--- {f.relative_to(out)} ---")
        print(f.read_text(encoding="utf-8-sig"))
    print(f"Self-test output folder: {out}")


def run_demo():
    """Dashboard with simulated instruments, saving to a throwaway demo folder."""
    import tempfile
    out_dir = Path(tempfile.mkdtemp(prefix="lablogger_demo_"))
    cfg = {"station_name": "Demo bench (simulated data)", "shared_file": {"enabled": True},
           "history": {"enabled": True, "keep_days": 0},
    "lot_rules": {"expire_minutes": 120, "lot_pattern": r"LOT-\d{4}", "lot_example": "LOT-1234",
                         "operators": ["Maria", "James", "Demo"]},
           "dashboard": {"port": 8080, "open_browser": True}, "upload": {"folder": str(out_dir / "Upload test")}}
    log(f"DEMO MODE: simulated instruments, data in {out_dir}")
    SHARED.configure(cfg, out_dir)
    configure_lot_rules(cfg)
    threading.Thread(target=context_watch, daemon=True).start()
    EXCEL.configure(cfg, out_dir)
    HISTORY.configure(cfg, out_dir)
    CONTEXT.update(lot="", operator="")
    mb = SerialInstrument({"name": "MB90", "type": "mb90", "port": "COM6"}, out_dir)
    ph = SerialInstrument({"name": "pH_meter", "type": "ph", "port": "COM3"}, out_dir)
    set_status("MB90", "USB · COM6 (simulated)", True)
    set_status("pH_meter", "USB · COM3 (simulated)", True)
    colo_dir = out_dir / "Colorimeter Exports"
    fw = FolderWatcher({"name": "LS177C", "type": "colorimeter_folder", "watch_folder": str(colo_dir)}, out_dir)
    fw.start()
    set_status("AQUALAB3", "SKALA · unreachable (no internet?)", False, type="skala_api", label="AQUALAB 3")
    start_dashboard(cfg, out_dir)

    def mb90_test(final):
        head = ["MOISTURE DETERMINATION", "Halogen Moisture Analyzer", "Type MB90", "SNR(Drying Unit) 1234567",
                "Method Name Chickpea MC", "Drying Prog Standard", "Drying Temp 105°C", "Switch Off A60(1mg/60s)",
                "Start Weight 3.104 g", "00:00 min 0.00%MC"]
        for line in head:
            mb.handle_line(line)
        for i in range(1, 13):
            if STOP.wait(1.5):
                return
            v = final * (1 - 0.82 ** (i * 1.4))
            mb.handle_line(f"{i // 2:02d}:{30 * (i % 2):02d} min {v:.2f} %MC")
        for line in ["Total Time. 06:00 min", f"End Result . {final:.2f} %MC", "1.Jan.15 15:35",
                     "---------------END----------------"]:
            mb.handle_line(line)

    def loop():
        n = 0
        while not STOP.is_set():
            n += 1
            mb90_test(4.6 + (n % 3) * 0.3)
            if STOP.wait(3):
                return
            ph.handle_line("AB33PH 1234567890 1.00")
            ph.handle_line(f"0 2026/10/05 10:0{n % 10} 00{n} pH")
            ph.handle_line(f"pH {6.35 + n % 4 * 0.03:.2f} pH {33.0 + n:.1f} mV")
            ph.handle_line("23.0 C ATC")
            ph.handle_line("98.4 % -3.1 mV")
            if STOP.wait(3):
                return
            ph.flush()
            (colo_dir / f"demo_{n}.csv").write_text(
                f"NO,Test Name,DateTime,L*,a*,b*\n{n},{CONTEXT['lot']} roast,{now_iso()[:16]},{44 + n % 3 * 0.4:.2f},10.05,24.11\n",
                encoding="utf-8")
            if STOP.wait(15):
                return

    threading.Thread(target=loop, daemon=True).start()
    if sys.stdin is not None and sys.stdout is not None:
        threading.Thread(target=console_loop, daemon=True).start()
    try:
        while not STOP.is_set():
            WRITER.flush()
            housekeeping()
            STOP.wait(2)
    except KeyboardInterrupt:
        STOP.set()


def skala_test():
    cfg = load_config()
    inst = next((i for i in cfg["instruments"] if i.get("type") == "skala_api"), None)
    if not inst:
        sys.exit("No instrument with type 'skala_api' in lab_logger_config.json")
    out_dir = expand(cfg["output_folder"])
    f = SkalaFetcher(inst, out_dir)
    tok = f.read_token()
    print(f"Token file: {f.token_file}  ({'found, ' + str(len(tok)) + ' characters' if tok else 'MISSING or empty'})")
    print(f"sbs_serial: {f.sbs_serial or 'MISSING'}")
    print("Asking SKALA for recent batches (read-only)...")
    try:
        n = f.poll_once(verbose=True)
    except SkalaAuthError as e:
        sys.exit(f"Token problem: {e}\nCopy a fresh credential_token into {f.token_file}")
    except Exception as e:
        sys.exit(f"Could not reach SKALA: {e}\nIs the laptop online?")
    print(f"OK. {n} new reading(s) saved to {monthly_csv(out_dir, f.name_)}")


def main():
    ap = argparse.ArgumentParser(description="Bench instrument data logger")
    ap.add_argument("--list-ports", action="store_true", help="list serial ports")
    ap.add_argument("--monitor", metavar="PORT", help="print raw data from a port")
    ap.add_argument("--baud", type=int, default=9600, help="baud rate for --monitor")
    ap.add_argument("--selftest", action="store_true", help="test parsers with sample data")
    ap.add_argument("--skala-test", action="store_true", help="check the SKALA token and list recent readings")
    ap.add_argument("--demo", action="store_true", help="open the dashboard with simulated instruments")
    a = ap.parse_args()
    if a.list_ports:
        list_ports()
    elif a.monitor:
        monitor(a.monitor, a.baud)
    elif a.selftest:
        selftest()
    elif a.skala_test:
        skala_test()
    elif a.demo:
        run_demo()
    else:
        run_logger()


if __name__ == "__main__":
    try:
        main()
    except StartupProblem as e:
        write_log_file(f"COULD NOT START: {e}")
        if sys.stdout is not None:
            print(f"\n  Lab Logger could not start:\n  {e}\n")
        sys.exit(2)
    except SystemExit:
        raise
    except BaseException:
        import traceback
        tb = traceback.format_exc()
        write_log_file("COULD NOT START (unexpected error):\n" + tb)
        if sys.stdout is not None:
            print(tb)
        sys.exit(3)
